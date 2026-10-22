"""推荐题库只读原任务、两次独立审阅和真实表格交付；不访问模型。"""
import copy
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile
from shared import question_recommendations as qr
from scripts import frontmind_workflow as wf


def fixture():
    return {'questions': [{'category': c, 'question': f'{c}测试问题第{i}项如何选择？'}
                          for c, count in qr.COUNTS.items() for i in range(count)], 'review_notes': ['独立复核完成']}


class RecommendationTests(unittest.TestCase):
    def test_counts_duplicates_and_no_fabricated_monitoring(self):
        self.assertEqual(qr.validate(fixture()), fixture())
        for change in ('missing', 'duplicate', 'answers', 'category'):
            value = fixture()
            if change == 'missing': value['questions'].pop()
            if change == 'duplicate': value['questions'][-1]['question'] = value['questions'][0]['question']
            if change == 'answers': value['answers'] = ['fabricated']
            if change == 'category': value['questions'][0]['category'] = '其他'
            with self.assertRaises(ValueError): qr.validate(value)

    def test_material_selection_preserves_full_text_and_hashes(self):
        with tempfile.TemporaryDirectory() as d:
            pack = Path(d)/'pack.zip'
            with zipfile.ZipFile(pack, 'w') as z:
                z.writestr('reference_pack.json', '{"brand":"测试品牌"}')
                z.writestr('strategy/confirmed.md', '完整定位原文')
                z.writestr('p0/article.md', '完整P0原文')
                z.writestr('provider/key.txt', 'must not enter context')
                z.writestr('strategy/../../private.txt', 'must not enter context')
            text = qr.pack_text(pack)
            self.assertIn('完整定位原文', text);self.assertIn('完整P0原文', text)
            self.assertIn('SHA256', text);self.assertNotIn('must not', text)

    def test_existing_job_is_unchanged_and_outputs_are_not_monitoring(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            job = Path(d)/'existing';job.mkdir()
            pack = Path(d)/'pack';(pack/'strategy').mkdir(parents=True)
            (pack/'strategy/confirmed.md').write_text('已确认的完整定位资料', encoding='utf-8')
            state = {'status':'p0_ready','job_id':'original-job','metadata':{'reference_pack_delivery':{'pack_path':str(pack)}}}
            raw = json.dumps(state, ensure_ascii=False).encode();(job/'job_state.json').write_bytes(raw)
            stub = SimpleNamespace(ROOT=Path(d), WorkflowError=wf.WorkflowError,
                assert_regular_source=lambda *a, **k: job, load_state=lambda *a: copy.deepcopy(state),
                validate_reference_pack=lambda *a: {'status':'pass','readiness':{'positioning_ready':True}}, atomic_json=wf.atomic_json)
            calls = []
            def action(root, original, name, prompt, validator, **kwargs):
                calls.append((original, name, prompt));return validator(fixture())
            with patch('shared.model_runtime.run_action', side_effect=action):
                self.assertEqual(qr.run(stub, SimpleNamespace(job_dir=str(job), retry_current_action=False)), 0)
            self.assertEqual([c[1] for c in calls], list(qr.ACTIONS))
            self.assertIn('待审题目', calls[1][2]);self.assertTrue(all(c[0] == job for c in calls))
            self.assertEqual((job/'job_state.json').read_bytes(), raw)
            output=job/'deliverables/question_recommendations'
            book=load_workbook(output/'recommended_questions.xlsx',read_only=True)
            self.assertEqual(book.active.max_row,161);book.close()
            self.assertEqual(json.loads((output/'recommended_questions.json').read_text())['source'], 'generated_recommendations_not_monitoring')

    def test_active_or_pending_job_cannot_launch_recommendation(self):
        for state in ({'status':'running_article'}, {'status':'p0_ready','pending_action':'existing'}):
            stub=SimpleNamespace(WorkflowError=wf.WorkflowError, assert_regular_source=lambda *a, **k: Path('.'), load_state=lambda *a: state)
            with patch('shared.model_runtime.run_action') as model, self.assertRaises(wf.WorkflowError):
                qr.run(stub,SimpleNamespace(job_dir='.', retry_current_action=False))
            model.assert_not_called()

    def test_host_requires_full_paginated_source_before_submission(self):
        from shared.host_tools import HostTools, ToolError
        with tempfile.TemporaryDirectory() as d:
            job=Path(d)/'existing';wf.ensure_new_job(job,'article',offline=True)
            source=job/'inputs/question_recommendation_source.md'
            source.parent.mkdir(parents=True,exist_ok=True)
            source.write_text('完整资料' * 10001,encoding='utf-8')
            tools=HostTools(wf.ROOT,job,qr.ACTIONS[0])
            record=next(x for x in tools.artifacts.values() if x['category']=='recommendation_source')
            with self.assertRaises(ToolError):tools.validate_complete_reads()
            tools._read({'artifact_id':record['artifact_id'],'offset':0,'limit':24000})
            with self.assertRaises(ToolError):tools.validate_complete_reads()
            tools._read({'artifact_id':record['artifact_id'],'offset':24000,'limit':24000})
            self.assertTrue(tools.validate_complete_reads())
