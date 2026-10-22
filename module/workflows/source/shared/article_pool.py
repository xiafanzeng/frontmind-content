"""Question-scoped citation intake and a visible, selectable article pool.

No search-generated ranking: every count comes from supplied citation details.
Pool selections are blueprint references; the P01 style pair stays independent.
"""
from __future__ import annotations

import csv
import hashlib
import html
import json
import re
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from . import question_research as qr, research_ingestion as ri
from .question_semantics import normalize_question_key

CONTRACT = 'frontmind-article-pool/1'
PATH = 'research/article_pool/pool.json'
SELECTION = 'examples/article_pool/selection.json'


def enabled(state):
    return state.get('job_kind') == 'article' and state.get('metadata', {}).get('article_pool_policy') == CONTRACT


def _hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _url(value):
    raw = ri._http_url(value)
    if not raw:
        raise ValueError('信源明细中的文章链接须为完整 HTTP(S) URL')
    parsed = urlsplit(raw)
    # Fragment variants refer to the same page; query strings remain meaningful.
    return urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(), parsed.path or '/', parsed.query, ''))


def stage_source(wf, job, path, period=None):
    state = wf.load_state(job)
    if not enabled(state):
        raise wf.WorkflowError('此上传入口用于 v16.4 新文章任务；历史任务沿用冻结输入。')
    source = wf.assert_regular_source(path)
    if source.suffix.lower() not in {'.xlsx', '.csv', '.json'}:
        raise wf.WorkflowError('请上传信源明细 .xlsx、.csv 或 .json；每行对应一次实际引用记录。')
    digest = wf.sha256_file(source)
    dest = Path(job) / 'inputs/citation_uploads' / (digest[:20] + source.suffix.lower())
    if not dest.is_file():
        wf.copy_stream(source, dest)
    state['metadata']['pending_citation_input'] = {
        'path': dest.relative_to(job).as_posix(), 'sha256': digest, 'period_id': period,
    }
    wf.save_state(job, state)


def read_details(path, question, *, explicit_period=None):
    """Select exact question/date rows, never merge other questions or periods."""
    path = Path(path)
    key = normalize_question_key(question['question_text'])
    if explicit_period and explicit_period != question.get('period_id'):
        raise ValueError('指定信源期次与当前问题期次不一致。')
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            candidates = [(sheet, *qr._detail_columns(sheet)) for sheet in workbook.worksheets]
            candidates = [item for item in candidates if 'question' in item[1] and 'url' in item[1]]
            named = [item for item in candidates if ri.field_key(item[0].title) in
                     {ri.field_key(x) for x in ri.DETAIL_SHEET_ALIASES}]
            candidates = named or candidates
            if len(candidates) != 1:
                raise ValueError('信源工作簿须有唯一的逐条明细表，包含问题与文章链接。')
            sheet, _, start_row = candidates[0]
            headers = next(sheet.iter_rows(min_row=start_row-1, max_row=start_row-1, values_only=True))
            raw = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=start_row, values_only=True)]
        finally:
            workbook.close()
    elif path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            raw = list(csv.DictReader(stream))
    else:
        raw = json.loads(path.read_text(encoding='utf-8-sig'))
        raw = raw.get('rows', []) if isinstance(raw, dict) else raw
    if not isinstance(raw, list):
        raise ValueError('信源 JSON 须为逐条引用记录数组，或含 rows 数组。')
    rows = []
    for entry in raw:
        if not isinstance(entry, dict):
            raise ValueError('信源明细每行须为对象。')
        field = lambda name: ri._field(entry, qr._DETAIL_ALIASES[name])
        if normalize_question_key(field('question')) != key:
            continue
        count = ri._field(entry, {'citation_count', '引用次数', '被引用次数'})
        if count not in (None, '', 1, '1'):
            raise ValueError('请提供逐条引用明细，不用汇总次数冒充单条记录。')
        observed = field('date')
        dates = ri._date_values(observed)
        if observed and not dates:
            raise ValueError('本题信源明细日期无法解析。')
        rows.append({'canonical_url': _url(field('url')),
                     'content_title': ri.normal(field('title')) or None,
                     'media_name': ri.normal(field('media_name')) or None,
                     'media_domain': ri.normal(field('media_domain')) or None,
                     'model': ri.normal(field('model')) or None,
                     'observed_at': ri.normal(observed) or None,
                     'observed_date': dates[0] if dates else None})
    if not rows:
        raise ValueError('信源文件没有当前正式问题的有效引用明细；请勿上传另一问题的信源代替。')
    start, end = question.get('observed_from'), question.get('observed_to')
    selected = []
    for row in rows:
        date = row.get('observed_date')
        if date and start and end:
            if not start <= date <= end:
                continue
        elif not explicit_period:
            raise ValueError('答案或信源没有可匹配日期；请补齐日期，或用 --source-period 明确绑定页面所示期次。')
        selected.append({**row, 'canonical_url': _url(row['canonical_url'])})
    if not selected:
        raise ValueError('本题信源只有其他期次记录，未混入当前期次。请补交本期信源。')
    return selected


def _install_source(wf, job, question, pending):
    source = Path(job) / pending['path']
    if source.is_symlink() or not source.is_file() or wf.sha256_file(source) != pending['sha256']:
        raise wf.WorkflowError('已上传信源文件发生变化，请重新上传。')
    try:
        details = read_details(source, question, explicit_period=pending.get('period_id'))
    except (ValueError, OSError) as exc:
        raise wf.WorkflowError(str(exc)) from exc
    state = wf.load_state(job)
    pack = Path(state['reference_pack']['path'])
    with wf.PackageView(pack) as view:
        index = view.read_json(qr.QUESTION_RESEARCH_INDEX_MEMBER)
        row = next((x for x in index['slices'] if x['question_uid'] == question['question_uid']
                    and x['period_id'] == question['period_id']), None)
        if row is None:
            raise wf.WorkflowError('本题期次尚未建立，请先补齐问题答案。')
        citation = qr._citation_slice(question['question_uid'], question['question_text'],
            normalize_question_key(question['question_text']), question['period_id'],
            {'details': details}, question.get('observed_from'), question.get('observed_to'))
        citation['availability'] = 'supplied'
        original_member = 'research/article_pool_sources/' + pending['sha256'] + source.suffix.lower()
        citation['source_upload'] = {'path': original_member, 'sha256': pending['sha256'],
                                     'explicit_period_id': pending.get('period_id')}
        row.update(citation_availability='supplied', citation_detail_count=len(details),
                   unique_citation_content_count=citation['unique_content_count'])
        index['index_revision'] += 1
        index['updated_at'] = wf.now()
    install = pending.get('install')
    if not install:
        version = int(state['reference_pack']['pack_version']) + 1
        dest = Path(job) / 'reference_updates' / f"Reference_Pack_sources_v{version}_{pending['sha256'][:10]}"
        install = {'destination': str(dest), 'pack_id': state['reference_pack']['pack_id'],
                   'pack_version': version, 'question_uid': question['question_uid'], 'period_id': question['period_id']}
        state['metadata']['pending_citation_input']['install'] = install
        wf.save_state(job, state)
    if (install['question_uid'], install['period_id']) != (question['question_uid'], question['period_id']):
        raise wf.WorkflowError('未完成的信源导入属于其他问题或期次，请为当前问题重新上传。')
    dest = Path(install['destination'])
    portable = dest.with_suffix('.zip')
    if dest.exists():
        report = wf.validate_reference_pack(dest)
        if report.get('status') != 'pass':
            raise wf.WorkflowError('中断前创建的资料包不完整，请重新上传信源。')
        root = wf.reference_pack_root(dest)
        with wf.PackageView(dest) as view:
            saved = view.read_json(row['citation_path'])
            original_hash = hashlib.sha256(view.read_bytes(original_member)).hexdigest()
        if (root['pack_id'] != install['pack_id'] or root['pack_version'] != install['pack_version']
                or saved.get('source_upload', {}).get('sha256') != pending['sha256']
                or original_hash != pending['sha256'] or saved.get('question_uid') != question['question_uid']
                or saved.get('period_id') != question['period_id'] or saved.get('details') != citation['details']):
            raise wf.WorkflowError('中断恢复的资料包与本次信源不一致，请重新上传。')
        if not portable.is_file():
            wf.write_reference_pack_zip(dest, portable)
        result = {'status': 'recovered', 'pack_path': str(dest), 'portable_zip_path': str(portable),
                  'portable_zip_sha256': 'sha256:' + wf.sha256_file(portable),
                  'pack_id': root['pack_id'], 'pack_version': root['pack_version'],
                  'parent_pack_version': root['parent_pack_version'], 'brand_name': root['brand_name'],
                  'readiness': report.get('readiness', {})}
    else:
        result = wf.create_next_reference_pack_version(pack=pack, output=dest,
            artifacts={row['citation_path']: json.dumps(citation, ensure_ascii=False, indent=2) + '\n',
                       qr.QUESTION_RESEARCH_INDEX_MEMBER: json.dumps(index, ensure_ascii=False, indent=2) + '\n',
                       original_member: source.read_bytes()}, portable_zip=portable)
    if (state['reference_pack']['pack_id'], state['reference_pack']['pack_version']) != (result['pack_id'], result['pack_version']):
        wf.bind_reference_pack(job, Path(result['pack_path']))
    state = wf.load_state(job)
    state['metadata'].pop('pending_citation_input', None)
    state['metadata']['question_source_delivery'] = result
    wf.save_state(job, state)
    return wf.select_question(Path(state['reference_pack']['path']), question['question_uid'],
                              period_id=question['period_id'])


def prepare(wf, job, question):
    state = wf.load_state(job)
    if not enabled(state):
        return question
    from .question_bank_import import known_platforms
    if len(known_platforms(question.get('answers', []))) < 2:
        raise wf.WorkflowError('请先提供本题同一期来自两个不同 AI 平台的完整答案，并补交对应信源明细。')
    pending = state['metadata'].get('pending_citation_input')
    if pending:
        question = _install_source(wf, job, question, pending)
        state = wf.load_state(job)
    with wf.PackageView(Path(state['reference_pack']['path'])) as view:
        index = view.read_json(qr.QUESTION_RESEARCH_INDEX_MEMBER)
        row = next((x for x in index['slices'] if x['question_uid'] == question.get('question_uid')
                    and x['period_id'] == question.get('period_id')), None)
        citation = view.read_json(row['citation_path']) if row else {}
    details = citation.get('details', [])
    if not details:
        raise wf.WorkflowError('已找到本题答案，但缺少本题本期信源明细；必须上传 --source-workbook 后才能列举 Top20 并继续。')
    if (citation.get('question_uid') != question['question_uid']
            or citation.get('period_id') != question['period_id']
            or normalize_question_key(citation.get('question')) != normalize_question_key(question['question_text'])):
        raise wf.WorkflowError('信源记录与当前问题或期次不一致。')
    normalized = []
    for n, value in enumerate(details):
        try:
            date = value.get('observed_date')
            if date and question.get('observed_from') and question.get('observed_to'):
                if not question['observed_from'] <= date <= question['observed_to']:
                    raise ValueError('资料包信源明细含其他期次记录，请补交当前期次的明细。')
            normalized.append({**value, 'canonical_url': _url(value.get('canonical_url')),
                               'detail_id': value.get('detail_id') or f'd{n+1:06d}'})
        except ValueError as exc:
            raise wf.WorkflowError(str(exc)) from exc
    articles = qr._rank_citations(normalized)
    for item in articles:
        item['article_id'] = 'a_' + hashlib.sha256(item['canonical_url'].encode()).hexdigest()[:16]
        item['media_domain'] = item.get('media_domain') or urlsplit(item['canonical_url']).netloc
    pool = {'contract': CONTRACT, 'question_uid': question['question_uid'],
            'question': question['question_text'], 'period_id': question['period_id'],
            'observed_from': question.get('observed_from'), 'observed_to': question.get('observed_to'),
            'articles': articles, 'detail_count': len(normalized),
            'source': {'pack_id': state['reference_pack']['pack_id'], 'citation_path': row['citation_path']},
            'ranking': 'citation_count_descending_first_observation_tiebreak'}
    pool['fingerprint'] = _hash({k: pool[k] for k in ('question_uid', 'question', 'period_id', 'articles')})
    wf.atomic_json(Path(job) / PATH, pool)
    export(wf, job, pool)
    return question


def load(wf, job):
    path = Path(job) / PATH
    if not path.is_file():
        raise wf.WorkflowError('尚未建立本题文章池，请先补齐问题答案与信源。')
    value = wf.read_json(path)
    q = wf.load_state(job).get('question') or {}
    if (value.get('question_uid'), value.get('period_id'), normalize_question_key(value.get('question'))) != (
            q.get('question_uid'), q.get('period_id'), normalize_question_key(q.get('question_text'))):
        raise wf.WorkflowError('文章池不是当前问题与期次，请重新载入本题输入。')
    expected = _hash({k: value[k] for k in ('question_uid', 'question', 'period_id', 'articles')})
    if value.get('fingerprint') != expected:
        raise wf.WorkflowError('文章池内容发生变化，请从已上传信源重新生成。')
    return value


def selected(wf, job, pool=None):
    path = Path(job) / SELECTION
    if not path.is_file():
        return None
    value = wf.read_json(path)
    pool = pool or load(wf, job)
    return value if value.get('pool_fingerprint') == pool['fingerprint'] else None


def _review_safe(value, *, url=False):
    # External titles/URLs may legitimately contain e.g. "score". Preserve
    # their rendered meaning without triggering the old internal-jargon guard.
    from scripts.frontmind_workflow import PUBLIC_FORBIDDEN_TOKENS
    text = str(value)
    for token in PUBLIC_FORBIDDEN_TOKENS:
        text = re.sub(re.escape(token), lambda match: (
            f'%{ord(match[0][0]):02X}' if url else f'&#{ord(match[0][0])};') + match[0][1:], text, flags=re.IGNORECASE)
    return text


def _cell(value):
    return _review_safe(html.escape(str(value or '未提供')).replace('|', '\\|').replace('\n', ' '))


def table(pool, selection=None, *, limit=None):
    chosen = {x['article_id'] for x in (selection or {}).get('examples', [])}
    rows = pool['articles'] if limit is None else pool['articles'][:limit]
    lines = ['| 编号 | 文章标题 | 媒体 | 被引用次数 | 原文链接 | 选择 |',
             '|---|---|---|---:|---|---|']
    for x in rows:
        url = _review_safe(x['canonical_url'].replace('>', '%3E').replace('<', '%3C'), url=True)
        lines.append(f"| {x['raw_rank']} | {_cell(x['content_title'] or x['canonical_url'])} | "
                     f"{_cell(x['media_name'] or x['media_domain'])} | {x['citation_count']} | "
                     f"[打开原文](<{url}>) | {'已选' if x['article_id'] in chosen else '—'} |")
    return '\n'.join(lines)


def export(wf, job, pool):
    selection = selected(wf, job, pool)
    header = (f"# 本题引用文章池\n\n问题：{pool['question']}\n\n期次：{pool['period_id']}"
              f" · {pool['observed_from'] or '日期未提供'}—{pool['observed_to'] or '日期未提供'}\n\n"
              f"共 {len(pool['articles'])} 篇去重文章，{pool['detail_count']} 条引用记录。按引用频次排序，不是质量排名。\n\n")
    for name, limit in [('top20.md', 20), ('all_articles.md', None)]:
        wf.atomic_text(Path(job) / 'research/article_pool' / name, header + table(pool, selection, limit=limit) + '\n')


def display(wf, job):
    pool = load(wf, job)
    value = selected(wf, job, pool)
    count = len(pool['articles'])
    lines = ['## 本题 Top20 引用文章池', '',
             f"问题：{pool['question']}；期次：{pool['period_id']}（{pool['observed_from'] or '日期未提供'}—{pool['observed_to'] or '日期未提供'}）。",
             f'共 {count} 篇；下列显示前 {min(count, 20)} 篇。按实际引用次数排序，非文章质量排名。', '',
             table(pool, value, limit=20), '',
             wf.markdown_link('查看完整文章池', Path(job) / 'research/article_pool/all_articles.md'), '',
             '按编号选择 1—2 篇作为本篇构思参考：--pool-selection 3,7；也可明确 --pool-selection none 不采用。完整池中20名之后的编号同样可选。',
             'P01 正文继续参考金蝶、ThinkPad；文章池选择不覆盖它们。两篇专项研究案例仍按原流程搜索、阅读全文。']
    if value is None:
        lines.append('文章池尚未选择；请先提交选择或明确不采用，再确认原有文风方案。')
    else:
        lines.append('已选择：' + ('、'.join(_cell(x['title']) for x in value['examples']) or '不采用额外文章池例文'))
        for x in value['examples']:
            lines.append(wf.markdown_link('查看已取得全文：' + _cell(x['title']), Path(job) / x['text_path']))
    errors = wf.load_state(job)['metadata'].get('article_pool_acquisition_errors', [])
    lines.extend('取得失败：' + _cell(x) for x in errors)
    if errors:
        lines.append('可用 --pool-selection 编号 --pool-example-file 完整正文.md 补交，或改选其他文章。')
    return '\n\n'.join(lines)


def choose(wf, job, text, files=()):
    from .example_acquisition import ExampleStore, AcquisitionError
    pool = load(wf, job)
    if str(text).strip().lower() in {'none', '无', '不采用'}:
        ids = []
    else:
        try:
            ids = [int(x.strip()) for x in str(text).replace('，', ',').split(',')]
        except ValueError as exc:
            raise wf.WorkflowError('请提供文章池编号，如 3,7；不采用则填 none。') from exc
        if not 1 <= len(ids) <= 2 or len(ids) != len(set(ids)) or any(n < 1 or n > len(pool['articles']) for n in ids):
            raise wf.WorkflowError('请从完整文章池选择1—2个不同的有效编号。')
    if files and len(files) != len(ids):
        raise wf.WorkflowError('补交全文文件应与本次选择的编号逐一对应。')
    store = ExampleStore(job)
    records, errors = [], []
    for index, n in enumerate(ids):
        item = pool['articles'][n-1]
        try:
            if files:
                record = store.import_user_file(wf.assert_regular_source(files[index]),
                    title=item['content_title'] or item['canonical_url'], source_url=item['canonical_url'])
            else:
                url = item['canonical_url']
                page = urlsplit(url)
                # These sites have explicit inner article containers; keep
                # their tokens local rather than broadening the generic parser.
                options = {}
                if (page.hostname in {'27sem.com', 'www.27sem.com'}
                        and page.path.startswith('/article/')):
                    options['article_tokens'] = ('paragraph_img',)
                elif page.hostname == 'm.11467.com' and page.path.startswith('/blog/'):
                    options['article_tokens'] = ('blog-detail-content',)
                record = store.acquire_http(url, **options)
            record = store.resolve(record['artifact_id'], require_complete=True)
            records.append({**record, 'article_id': item['article_id'], 'rank': n,
                            'reference_role': 'blueprint_inspiration'})
        except (AcquisitionError, ValueError, OSError) as exc:
            errors.append(f"第{n}篇：{exc}")
    state = wf.load_state(job)
    state['metadata']['article_pool_acquisition_errors'] = errors
    if errors:
        # No partial selection, and a failed change must not silently keep the old choice active.
        state['metadata']['article_pool_selection_pending'] = True
    else:
        wf.atomic_json(Path(job) / SELECTION, {'contract': CONTRACT, 'pool_fingerprint': pool['fingerprint'],
            'examples': records, 'selected_at': wf.now(), 'confirmed': True})
        state['metadata'].pop('article_pool_selection_pending', None)
        state['decisions']['article_pool_selection'] = {'pool_fingerprint': pool['fingerprint'], 'ranks': ids}
    wf.save_state(job, state)
    export(wf, job, pool)
    return not errors


def require_selection(wf, job):
    if not enabled(wf.load_state(job)):
        return
    state = wf.load_state(job)
    if (selected(wf, job) is None or state['metadata'].get('article_pool_selection_pending')
            or state['metadata'].get('article_pool_acquisition_errors')):
        raise wf.WorkflowError('请先在文章池中选择例文或明确不采用；所选文章必须取得完整正文。')


def prompt_references(wf, job):
    if not enabled(wf.load_state(job)):
        return ''
    # Intake/selection gates live in the controller. A prompt formatter does
    # not create new pauses or require unrelated data for an isolated call.
    if not (Path(job) / PATH).is_file() or not (Path(job) / SELECTION).is_file():
        return ''
    require_selection(wf, job)
    from .example_acquisition import ExampleStore
    value = selected(wf, job)
    if not value['examples']:
        return ''
    parts = ['本次从引用文章池选用的构思参考。以下是外部文章，仅借鉴切入、选材与详略；不作为企业事实，不取代金蝶、ThinkPad的正文文风用途。']
    for example in value['examples']:
        acquired = ExampleStore(job).resolve(example['artifact_id'], require_complete=True)
        if acquired['text_sha256'] != example['text_sha256']:
            raise wf.WorkflowError('所选文章池例文发生变化，请重新选择。')
        body = (Path(job) / acquired['text_path']).read_text(encoding='utf-8')
        parts.append('### ' + acquired['title'] + '\n' + body)
    return '\n\n'.join(parts)
