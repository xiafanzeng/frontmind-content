#!/usr/bin/env python3
"""Neutral ingestion primitives for question monitoring and citation exports.

These exports are execution context, never factual evidence or editorial
strategy.  The module inventories exact questions, dates, answer observations
and citations without deciding article structure, presentation or claims.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import stat
import tempfile
import unicodedata
import zipfile
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Protocol, Sequence
from urllib.parse import urlsplit

SCHEMA_VERSION = "1.0.0"
ARTIFACT_TYPE = "frontmind_question_research_ingestion_index"
USAGE_POLICY = "context_only_not_factual_evidence"
INDEX_MEMBER = "research/observations/index.json"
STAGING_INDEX_NAME = "index.json"
MAX_FILE_BYTES = 128 * 1024 * 1024
MAX_TOTAL_XLSX_BYTES = 512 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_PACKAGE_FILES = 5_000
MONITORING_SUFFIXES = {".xlsx", ".csv", ".json"}

QUESTION_ALIASES = {
    "question", "query", "prompt", "monitoringquestion", "monitorquestion",
    "问题", "监控问题", "正式问题",
}
PLATFORM_ALIASES = {"platform", "provider", "channel", "平台", "模型平台"}
DATE_ALIASES = {
    "date", "datetime", "time", "capturedat", "sampledat", "sampled_at",
    "日期", "时间", "监控时间", "引用日期",
}
ANSWER_ALIASES = {
    "answer", "answertext", "answer_text", "content", "response", "output",
    "答案", "答案内容", "内容", "回答",
}
BRAND_ALIASES = {"brand", "brandname", "company", "companyname", "品牌", "品牌名称", "企业", "企业名称"}
URL_ALIASES = {"url", "contenturl", "内容url", "内容链接", "链接", "文章链接"}
DETAIL_SHEET_ALIASES = {"详细表格", "引用明细", "信源明细", "details", "detailedtable"}
RANK_ALIASES = {"rank", "ranking", "position", "排名", "名次", "监控词排名"}
THEME_LEXICON = {
    "产品服务与核心能力": ("产品", "服务", "功能", "课程", "方案", "项目", "能力", "内容"),
    "团队资质与专业基础": ("团队", "人员", "经验", "资质", "认证", "师资", "专家", "医生", "医师"),
    "需求适配与使用条件": ("适合", "适用", "需求", "场景", "条件", "人群", "对象", "个体化"),
    "价格成本与交易条件": ("价格", "费用", "成本", "收费", "预算", "报价", "套餐", "合同"),
    "风险限制与注意事项": ("风险", "限制", "安全", "缺点", "不足", "副作用", "禁忌", "注意事项"),
    "实施交付与使用流程": ("流程", "实施", "部署", "交付", "安装", "预约", "办理", "咨询", "使用"),
    "支持维护与持续服务": ("支持", "维护", "售后", "更新", "随访", "复诊", "护理", "服务"),
    "体验口碑与案例证据": ("体验", "口碑", "评价", "投诉", "案例", "效果", "结果", "改善", "疗效"),
}


class ResearchIngestionError(ValueError):
    """A monitoring or citation research input cannot be used safely."""

    def __init__(self, message: str, *, excluded_members: Sequence[str] = ()):
        super().__init__(message)
        self.excluded_members = tuple(dict.fromkeys(str(item) for item in excluded_members if str(item)))


class PackageReader(Protocol):
    def names(self) -> list[str]: ...
    def read(self, name: str) -> bytes: ...


class _DirectoryPackageReader:
    def __init__(self, root: Path):
        self.root = root.resolve()
        if not self.root.is_dir():
            raise ResearchIngestionError(f"not a Reference Pack directory: {root}")
        for path in self.root.rglob("*"):
            if path.is_symlink():
                raise ResearchIngestionError(f"Reference Pack directory contains a symlink: {path.relative_to(self.root)}")

    def names(self) -> list[str]:
        names = sorted(path.relative_to(self.root).as_posix() for path in self.root.rglob("*") if path.is_file())
        if len(names) > MAX_PACKAGE_FILES:
            raise ResearchIngestionError(f"Reference Pack contains more than {MAX_PACKAGE_FILES} files")
        total = 0
        for name in names:
            if not safe_relative(name):
                raise ResearchIngestionError(f"unsafe Reference Pack member: {name}")
            size = (self.root / name).stat().st_size
            if size > MAX_FILE_BYTES:
                raise ResearchIngestionError(f"Reference Pack member is too large: {name}")
            total += size
        if total > MAX_TOTAL_XLSX_BYTES:
            raise ResearchIngestionError("Reference Pack exceeds the uncompressed size limit")
        return names

    def read(self, name: str) -> bytes:
        if not safe_relative(name):
            raise ResearchIngestionError(f"unsafe Reference Pack member: {name}")
        return (self.root / name).read_bytes()

    def close(self) -> None:
        return None


class _ZipPackageReader:
    def __init__(self, path: Path):
        try:
            self.archive = zipfile.ZipFile(path)
        except (OSError, zipfile.BadZipFile) as exc:
            raise ResearchIngestionError(f"unreadable Reference Pack ZIP: {path}") from exc
        infos = [item for item in self.archive.infolist() if not item.is_dir()]
        names = [item.filename for item in infos]
        if len(names) != len(set(names)):
            raise ResearchIngestionError("Reference Pack ZIP contains duplicate paths")
        if len(names) > MAX_PACKAGE_FILES:
            raise ResearchIngestionError(f"Reference Pack ZIP contains more than {MAX_PACKAGE_FILES} files")
        total = 0
        for item in infos:
            if not safe_relative(item.filename):
                raise ResearchIngestionError(f"unsafe Reference Pack ZIP member: {item.filename}")
            if item.flag_bits & 0x1:
                raise ResearchIngestionError(f"encrypted Reference Pack member is unsupported: {item.filename}")
            mode = item.external_attr >> 16
            if mode and stat.S_ISLNK(mode):
                raise ResearchIngestionError(f"Reference Pack ZIP contains a symlink: {item.filename}")
            if item.file_size > MAX_FILE_BYTES:
                raise ResearchIngestionError(f"Reference Pack member is too large: {item.filename}")
            if item.compress_size and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO:
                raise ResearchIngestionError(f"Reference Pack member exceeds the compression-ratio limit: {item.filename}")
            total += item.file_size
        if total > MAX_TOTAL_XLSX_BYTES:
            raise ResearchIngestionError("Reference Pack ZIP exceeds the uncompressed size limit")

    def names(self) -> list[str]:
        return sorted(item.filename for item in self.archive.infolist() if not item.is_dir())

    def read(self, name: str) -> bytes:
        if not safe_relative(name):
            raise ResearchIngestionError(f"unsafe Reference Pack member: {name}")
        try:
            return self.archive.read(name)
        except KeyError as exc:
            raise ResearchIngestionError(f"missing Reference Pack member: {name}") from exc

    def close(self) -> None:
        self.archive.close()


def normal(value: Any) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or ""))).strip()


def field_key(value: Any) -> str:
    return re.sub(
        r"[^0-9a-z\u3400-\u9fff]+",
        "",
        unicodedata.normalize("NFKC", normal(value)).casefold(),
    )


def question_key(value: Any) -> str:
    """Exact normalized key shared with E2; it never performs fuzzy matching."""

    return field_key(value)


def safe_relative(value: Any) -> bool:
    if not isinstance(value, str) or not value or "\\" in value or "\x00" in value:
        return False
    path = PurePosixPath(value)
    return not path.is_absolute() and ".." not in path.parts and path.as_posix() == value


def _strict_json_bytes(data: bytes, label: str) -> Any:
    def reject_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for name, value in pairs:
            if name in result:
                raise ValueError(f"duplicate object key {name!r}")
            result[name] = value
        return result

    def reject_constant(value: str) -> None:
        raise ValueError(f"non-standard number {value}")

    try:
        return json.loads(
            data.decode("utf-8"),
            object_pairs_hook=reject_pairs,
            parse_constant=reject_constant,
        )
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError) as exc:
        raise ResearchIngestionError(f"invalid JSON {label}: {exc}") from exc


def _atomic_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def _regular_file(path: Path, suffixes: set[str], label: str) -> Path:
    candidate = path.expanduser().resolve()
    if path.is_symlink() or not candidate.is_file():
        raise ResearchIngestionError(f"{label} must be a regular non-symlink file: {path}")
    if candidate.suffix.casefold() not in suffixes:
        expected = ", ".join(sorted(suffixes))
        raise ResearchIngestionError(f"{label} must use one of {expected}: {candidate.name}")
    if candidate.stat().st_size > MAX_FILE_BYTES:
        raise ResearchIngestionError(f"{label} exceeds the {MAX_FILE_BYTES}-byte limit: {candidate.name}")
    if candidate.suffix.casefold() == ".xlsx":
        _validate_xlsx_container(candidate, label)
    return candidate


def _validate_xlsx_container(path: Path, label: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = [item for item in archive.infolist() if not item.is_dir()]
            names = [item.filename for item in infos]
            if len(names) != len(set(names)):
                raise ResearchIngestionError(f"{label} XLSX contains duplicate members")
            total = 0
            for item in infos:
                if not safe_relative(item.filename):
                    raise ResearchIngestionError(f"{label} XLSX contains an unsafe member: {item.filename}")
                if item.flag_bits & 0x1:
                    raise ResearchIngestionError(f"{label} XLSX contains an encrypted member: {item.filename}")
                mode = item.external_attr >> 16
                if mode and stat.S_ISLNK(mode):
                    raise ResearchIngestionError(f"{label} XLSX contains a symbolic link: {item.filename}")
                if item.file_size > MAX_FILE_BYTES:
                    raise ResearchIngestionError(f"{label} XLSX member is too large: {item.filename}")
                if item.compress_size and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO:
                    raise ResearchIngestionError(f"{label} XLSX member exceeds the compression-ratio limit: {item.filename}")
                total += item.file_size
            if total > MAX_TOTAL_XLSX_BYTES:
                raise ResearchIngestionError(f"{label} XLSX exceeds the uncompressed size limit")
    except zipfile.BadZipFile as exc:
        raise ResearchIngestionError(f"{label} is not a readable XLSX file: {path.name}") from exc


def _date_value(value: Any) -> str | None:
    values = _date_values(value)
    return values[0] if values else None


def _date_values(value: Any) -> list[str]:
    if isinstance(value, datetime):
        return [value.date().isoformat()]
    if isinstance(value, date):
        return [value.isoformat()]
    text = normal(value)
    if not text:
        return []
    result: list[str] = []
    for match in re.finditer(r"(?<!\d)(20\d{2})[年./-](\d{1,2})[月./-](\d{1,2})", text):
        try:
            result.append(date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat())
        except ValueError:
            continue
    return list(dict.fromkeys(result))


def _http_url(value: Any) -> str | None:
    raw = normal(value)
    try:
        parsed = urlsplit(raw)
    except ValueError:
        return None
    return raw if parsed.scheme.casefold() in {"http", "https"} and parsed.netloc else None


def _field(row: dict[str, Any], aliases: set[str]) -> Any:
    for name, value in row.items():
        if field_key(name) in {field_key(alias) for alias in aliases} and value not in (None, ""):
            return value
    return None


def _answer_values(row: dict[str, Any]) -> list[Any]:
    result: list[Any] = []
    normalized_aliases = {field_key(alias) for alias in ANSWER_ALIASES}
    for name, value in row.items():
        if value in (None, ""):
            continue
        name_key = field_key(name)
        if name_key in normalized_aliases or (
            (name_key.startswith("answer") or name_key.startswith("答案"))
            and (
                name_key.endswith("content")
                or name_key.endswith("内容")
                or re.fullmatch(r"(?:answer|答案)\d*", name_key)
            )
        ) or (
            # The real monitoring export uses a two-row ``品牌 / 内容``
            # header, composed as ``品牌_内容``.  Treat that terminal role as
            # answer text while excluding links/screenshots and metadata.
            (name_key.endswith("内容") or name_key.endswith("content"))
            and not any(token in name_key for token in ("url", "链接", "截图", "问题", "question"))
            and _http_url(value) is None
        ):
            result.append(value)
    return result


def _brand_needles(canonical_brand: str, aliases: Sequence[str]) -> tuple[str, ...]:
    values: list[str] = []
    for value in (canonical_brand, *aliases):
        text = unicodedata.normalize("NFKC", normal(value)).casefold()
        if len(question_key(text)) >= 2:
            values.append(text)
        without_qualifier = re.sub(r"[（(][^()（）]{1,80}[)）]", "", text).strip()
        if len(question_key(without_qualifier)) >= 2:
            values.append(without_qualifier)
    return tuple(dict.fromkeys(values))


def _mentions_brand(value: Any, needles: Sequence[str]) -> bool:
    text = unicodedata.normalize("NFKC", normal(value)).casefold()
    return bool(text and any(needle in text for needle in needles))


def _numeric_rank(value: Any) -> float | None:
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        match = re.search(r"(?<!\d)(\d+(?:\.\d+)?)(?!\d)", normal(value))
        if not match:
            return None
        number = float(match.group(1))
    return number if 0 < number <= 1_000 else None


def _rank_values(row: dict[str, Any], needles: Sequence[str], row_brand: str) -> list[float]:
    ranks: list[float] = []
    brand_key_values = {question_key(value) for value in needles if question_key(value)}
    row_matches = _mentions_brand(row_brand, needles) if row_brand else False
    for name, value in row.items():
        name_key = field_key(name)
        if not any(field_key(alias) in name_key for alias in RANK_ALIASES):
            continue
        # A brand-specific composed header is reliable.  A generic rank field
        # is accepted only when the row itself declares the canonical brand.
        brand_specific = any(candidate and candidate in name_key for candidate in brand_key_values)
        generic = name_key in {field_key(alias) for alias in RANK_ALIASES}
        if not brand_specific and not (generic and row_matches):
            continue
        rank = _numeric_rank(value)
        if rank is not None:
            ranks.append(rank)
    return ranks


def _theme_hits(value: Any) -> set[str]:
    text = normal(value)
    return {
        theme for theme, tokens in THEME_LEXICON.items()
        if any(token in text for token in tokens)
    }


def _new_monitor_metrics_state(needles: Sequence[str]) -> dict[str, Any]:
    return {
        "brand_metric_available": bool(needles),
        "answer_count": 0,
        "brand_mention_count": 0,
        "ranks": [],
        "platforms": {},
        "themes": Counter(),
    }


def _accumulate_monitor_metrics(
    state: dict[str, Any],
    answers: Sequence[Any],
    ranks: Sequence[float],
    platform: str,
    needles: Sequence[str],
) -> None:
    platform_state = state["platforms"].setdefault(
        platform or "unspecified",
        {"answer_count": 0, "brand_mention_count": 0, "ranks": []},
    )
    for answer in answers:
        state["answer_count"] += 1
        platform_state["answer_count"] += 1
        if needles and _mentions_brand(answer, needles):
            state["brand_mention_count"] += 1
            platform_state["brand_mention_count"] += 1
        for theme in _theme_hits(answer):
            state["themes"][theme] += 1
    state["ranks"].extend(ranks)
    platform_state["ranks"].extend(ranks)


def _rate(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 6) if denominator else None


def _average(values: Sequence[float]) -> float | None:
    return round(sum(values) / len(values), 6) if values else None


def _finalize_monitor_metrics(state: dict[str, Any]) -> dict[str, Any]:
    answer_count = int(state["answer_count"])
    mention_count = int(state["brand_mention_count"])
    ranks = [float(item) for item in state["ranks"]]
    metric_available = bool(state["brand_metric_available"] and answer_count)
    warnings: list[str] = []
    if not metric_available:
        warnings.append("brand_mention_rate_unavailable")
    if not ranks:
        warnings.append("brand_rank_metrics_unavailable")
    platform_differences: list[dict[str, Any]] = []
    for platform, values in sorted(state["platforms"].items()):
        platform_answers = int(values["answer_count"])
        platform_mentions = int(values["brand_mention_count"])
        platform_ranks = [float(item) for item in values["ranks"]]
        platform_differences.append({
            "platform": platform,
            "answer_observation_count": platform_answers,
            "brand_mention_rate": _rate(platform_mentions, platform_answers) if metric_available else None,
            "rank_observation_count": len(platform_ranks),
            "average_rank": _average(platform_ranks),
        })
    repeated_themes = [
        {"theme": theme, "observation_count": count}
        for theme, count in sorted(state["themes"].items(), key=lambda item: (-item[1], item[0]))
        if count >= 2
    ]
    return {
        "brand_visibility": {
            "answer_observation_count": answer_count,
            "brand_mention_count": mention_count if metric_available else None,
            "brand_mention_rate": _rate(mention_count, answer_count) if metric_available else None,
            "rank_observation_count": len(ranks),
            "top1_observation_count": sum(rank <= 1 for rank in ranks),
            "top3_observation_count": sum(rank <= 3 for rank in ranks),
            "top1_rate": _rate(sum(rank <= 1 for rank in ranks), len(ranks)),
            "top3_rate": _rate(sum(rank <= 3 for rank in ranks), len(ranks)),
            "average_rank": _average(ranks),
            "platform_differences": platform_differences,
        },
        "repeated_themes": repeated_themes,
        "warnings": warnings,
    }


def _composed_headers(first: Sequence[Any], second: Sequence[Any] | None) -> tuple[list[str], bool]:
    first_names = [normal(item) for item in first]
    if second is None:
        return [name or f"column_{index + 1}" for index, name in enumerate(first_names)], False
    second_names = [normal(item) for item in second]
    two_rows = any(field_key(name) in {"内容", "content", "截图链接", "排名", "监控词排名"} for name in second_names)
    if not two_rows:
        return [name or f"column_{index + 1}" for index, name in enumerate(first_names)], False
    headers: list[str] = []
    for index in range(max(len(first_names), len(second_names))):
        top = first_names[index] if index < len(first_names) else ""
        bottom = second_names[index] if index < len(second_names) else ""
        headers.append(f"{top}_{bottom}" if top and bottom and field_key(top) != field_key(bottom) else top or bottom or f"column_{index + 1}")
    return headers, True


def _rows_from_xlsx(path: Path) -> list[tuple[str, dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ResearchIngestionError("openpyxl is required to inspect research-observation XLSX files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    result: list[tuple[str, dict[str, Any]]] = []
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, ())
            second = next(iterator, None)
            headers, consumed_second = _composed_headers(first, second)
            values_iter: Iterable[Sequence[Any]]
            if not consumed_second and second is not None:
                values_iter = [second, *iterator]
            else:
                values_iter = iterator
            for values in values_iter:
                if not any(item not in (None, "") for item in values):
                    continue
                result.append((normal(sheet.title) or path.stem, {
                    headers[index]: value for index, value in enumerate(values) if index < len(headers)
                }))
    finally:
        workbook.close()
    return result


def _rows_from_csv(path: Path) -> list[tuple[str, dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        return [("unspecified", dict(row)) for row in csv.DictReader(handle)]


def _monitoring_inventory(
    path: Path,
    canonical_brand: str = "",
    aliases: Sequence[str] = (),
) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    normalized_context: dict[str, Any] | None = None
    if suffix == ".xlsx":
        rows = _rows_from_xlsx(path)
    elif suffix == ".csv":
        rows = _rows_from_csv(path)
    elif suffix == ".json":
        value = _strict_json_bytes(path.read_bytes(), path.name)
        if isinstance(value, dict) and isinstance(value.get("answers"), list) and normal(value.get("question")):
            normalized_context = value
            rows = []
        else:
            raw_rows = value if isinstance(value, list) else value.get("rows") if isinstance(value, dict) else None
            if not isinstance(raw_rows, list):
                raise ResearchIngestionError("monitoring JSON must be a normalized context, an array, or an object with rows[]")
            rows = [("unspecified", item) for item in raw_rows if isinstance(item, dict)]
    else:  # pragma: no cover - checked by _regular_file
        raise ResearchIngestionError(f"unsupported monitoring file: {path.name}")

    questions: dict[str, dict[str, Any]] = {}
    brands: set[str] = set()
    all_dates: list[str] = []
    brand_needles = _brand_needles(canonical_brand, aliases)
    if normalized_context is not None:
        text = normal(normalized_context.get("question"))
        key = question_key(text)
        if not key:
            raise ResearchIngestionError("monitoring context has no identifiable question")
        answers = [item for item in normalized_context.get("answers", []) if isinstance(item, dict) and normal(item.get("answer_text"))]
        platforms = {normal(item.get("platform")) for item in answers if normal(item.get("platform"))}
        dates = [item for item in (_date_value(item.get("sampled_at")) for item in answers) if item]
        captured = _date_value(normalized_context.get("captured_at"))
        if captured:
            dates.append(captured)
        all_dates.extend(dates)
        metrics_state = _new_monitor_metrics_state(brand_needles)
        for answer in answers:
            platform = normal(answer.get("platform")) or "unspecified"
            row_brand = normal(answer.get("brand") or answer.get("brand_name"))
            ranks = _rank_values(answer, brand_needles, row_brand)
            _accumulate_monitor_metrics(
                metrics_state,
                [answer.get("answer_text")],
                ranks,
                platform,
                brand_needles,
            )
        questions[key] = {
            "texts": [text],
            "answer_observation_count": len(answers),
            "platforms": sorted(platforms),
            "dates": dates,
            "research_metrics": _finalize_monitor_metrics(metrics_state),
        }
    else:
        for sheet_name, row in rows:
            question = normal(_field(row, QUESTION_ALIASES))
            key = question_key(question)
            if not key:
                continue
            answers = _answer_values(row)
            if not answers:
                continue
            item = questions.setdefault(key, {
                "texts": [], "answer_observation_count": 0, "platforms": set(), "dates": [],
                "metrics_state": _new_monitor_metrics_state(brand_needles),
            })
            if question not in item["texts"]:
                item["texts"].append(question)
            item["answer_observation_count"] += len(answers)
            platform = normal(_field(row, PLATFORM_ALIASES)) or normal(sheet_name)
            if platform:
                item["platforms"].add(platform)
            observed_dates = _date_values(_field(row, DATE_ALIASES))
            if observed_dates:
                item["dates"].extend(observed_dates)
                all_dates.extend(observed_dates)
            brand = normal(_field(row, BRAND_ALIASES))
            if brand:
                brands.add(brand)
            _accumulate_monitor_metrics(
                item["metrics_state"],
                answers,
                _rank_values(row, brand_needles, brand),
                platform,
                brand_needles,
            )
        for item in questions.values():
            item["platforms"] = sorted(item["platforms"])
            item["research_metrics"] = _finalize_monitor_metrics(item.pop("metrics_state"))

    if not questions:
        raise ResearchIngestionError("monitoring file has no row with both an identifiable question and answer content")
    return {
        "questions": questions,
        "platforms": sorted({platform for item in questions.values() for platform in item["platforms"]}),
        "observed_from": min(all_dates) if all_dates else None,
        "observed_to": max(all_dates) if all_dates else None,
        "brands": sorted(brands),
    }


def _sheet_headers(sheet: Any) -> tuple[dict[str, int], int]:
    aliases = {
        "question": {field_key(item) for item in QUESTION_ALIASES},
        "url": {field_key(item) for item in URL_ALIASES},
        "date": {field_key(item) for item in DATE_ALIASES},
        "brand": {field_key(item) for item in BRAND_ALIASES},
    }
    for row_number in (1, 2):
        values = next(sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True), ())
        normalized = [field_key(value) for value in values]
        columns: dict[str, int] = {}
        for name, choices in aliases.items():
            for index, value in enumerate(normalized):
                if value in choices:
                    columns[name] = index
                    break
        if "question" in columns and "url" in columns:
            return columns, row_number + 1
    return {}, 2


def _source_inventory(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ResearchIngestionError("openpyxl is required to inspect the research-observation source workbook") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        detail_candidates: list[tuple[Any, dict[str, int], int]] = []
        alias_keys = {field_key(item) for item in DETAIL_SHEET_ALIASES}
        for sheet in workbook.worksheets:
            columns, start = _sheet_headers(sheet)
            if "question" in columns and "url" in columns:
                detail_candidates.append((sheet, columns, start))
        named = [item for item in detail_candidates if field_key(item[0].title) in alias_keys]
        candidates = named or detail_candidates
        if len(candidates) != 1:
            if not candidates:
                raise ResearchIngestionError("source workbook has no recognizable detail sheet with question and URL columns")
            raise ResearchIngestionError("source workbook has multiple possible detail sheets and no unique recognized detail role")
        sheet, columns, start = candidates[0]
        questions: dict[str, dict[str, Any]] = {}
        brands: set[str] = set()
        all_dates: list[str] = []
        for values in sheet.iter_rows(min_row=start, values_only=True):
            def cell(name: str) -> Any:
                index = columns.get(name)
                return values[index] if index is not None and index < len(values) else None

            question = normal(cell("question"))
            key = question_key(question)
            url = _http_url(cell("url"))
            if not key or not url:
                continue
            item = questions.setdefault(key, {
                "texts": [], "citation_detail_count": 0, "dates": [], "domains": Counter(),
            })
            if question not in item["texts"]:
                item["texts"].append(question)
            item["citation_detail_count"] += 1
            hostname = (urlsplit(url).hostname or "").casefold().removeprefix("www.")
            if hostname:
                item["domains"][hostname] += 1
            observed_dates = _date_values(cell("date"))
            if observed_dates:
                item["dates"].extend(observed_dates)
                all_dates.extend(observed_dates)
            brand = normal(cell("brand"))
            if brand:
                brands.add(brand)
        for item in questions.values():
            total = int(item["citation_detail_count"])
            ordered_domains = sorted(item.pop("domains").items(), key=lambda value: (-value[1], value[0]))
            item["research_metrics"] = {
                "citation_domain_count": len(ordered_domains),
                "top_citation_domains": [
                    {
                        "domain": domain,
                        "citation_count": count,
                        "citation_share": _rate(count, total),
                    }
                    for domain, count in ordered_domains[:10]
                ],
                "citation_domain_concentration_top3": _rate(
                    sum(count for _domain, count in ordered_domains[:3]), total,
                ),
                "warnings": ["citation_domains_unavailable"] if not ordered_domains else [],
            }
        return {
            "questions": questions,
            "observed_from": min(all_dates) if all_dates else None,
            "observed_to": max(all_dates) if all_dates else None,
            "brands": sorted(brands),
            "detail_sheet": sheet.title,
        }
    finally:
        workbook.close()


def _question_records(project_questions: Any) -> dict[str, dict[str, Any]]:
    if isinstance(project_questions, dict):
        values = project_questions.get("questions", [])
    else:
        values = project_questions or []
    result: dict[str, dict[str, Any]] = {}
    for order, item in enumerate(values, 1):
        if isinstance(item, str):
            text = normal(item)
            question_id = None
            source_kind = "project_optimization_target"
        elif isinstance(item, dict):
            text = normal(item.get("question_text") or item.get("question") or item.get("text"))
            question_id = normal(item.get("question_id")) or None
            source_kind = normal(item.get("source_kind")) or "project_optimization_target"
        else:
            continue
        if source_kind != "project_optimization_target" or not question_key(text):
            continue
        result.setdefault(question_key(text), {"question_id": question_id, "question_text": text, "source_order": order})
    return result


def _brand_key(value: Any) -> str:
    return question_key(re.sub(r"[（(][^()（）]{1,80}[)）]", "", normal(value)))


def _brand_compatible(observed: Sequence[str], canonical_brand: str, aliases: Sequence[str]) -> tuple[bool, str | None]:
    if not observed:
        return True, None
    accepted = {_brand_key(value) for value in (canonical_brand, *aliases) if _brand_key(value)}
    for raw in observed:
        candidate = _brand_key(raw)
        if candidate in accepted:
            return True, None
        if len(candidate) >= 4 and any(value.startswith(candidate) or candidate.startswith(value) for value in accepted):
            return True, f"research-observation brand {raw!r} matched the canonical brand through a parent/alias normalization"
    return False, f"research observation declares unrelated brand values: {' | '.join(observed[:5])}"


def _effective_dates(monitoring: dict[str, Any], source: dict[str, Any]) -> tuple[str | None, str | None, list[str]]:
    monitor_from, monitor_to = monitoring.get("observed_from"), monitoring.get("observed_to")
    source_from, source_to = source.get("observed_from"), source.get("observed_to")
    warnings: list[str] = []
    if monitor_from and source_from:
        start = max(monitor_from, source_from)
        end = min(monitor_to, source_to)
        if start > end:
            raise ResearchIngestionError("monitoring answers and source workbook have disjoint observed date ranges")
        return start, end, warnings
    if monitor_from or source_from:
        warnings.append("only one side of the research observation pair has a usable observed date range")
        return monitor_from or source_from, monitor_to or source_to, warnings
    warnings.append("the research observation pair has no usable observed dates; ingested_at will be the selection fallback")
    return None, None, warnings


def _build_dataset(
    monitoring: dict[str, Any],
    source: dict[str, Any],
    *,
    dataset_id: str,
    content_fingerprint: str,
    monitoring_member: str,
    source_member: str,
    project_questions: Any,
    ingested_at: str,
    declared_period_label: str | None,
    canonical_brand: str,
    aliases: Sequence[str],
) -> dict[str, Any]:
    monitor_keys = set(monitoring["questions"])
    source_keys = set(source["questions"])
    citation_placeholders: list[str] = []
    for key in sorted(monitor_keys - source_keys):
        monitor = monitoring["questions"][key]
        source["questions"][key] = {
            "texts": list(monitor.get("texts") or []),
            "citation_detail_count": 0,
            "research_metrics": {
                "citation_domain_count": 0,
                "top_citation_domains": [],
                "citation_domain_concentration_top3": None,
                "warnings": ["citation_details_unavailable"],
            },
        }
        citation_placeholders.append(key)
    source_keys = set(source["questions"])
    if not monitor_keys:
        raise ResearchIngestionError("monitoring answers contain no usable question")
    observed_brands = sorted(set(monitoring.get("brands", [])) | set(source.get("brands", [])))
    brand_ok, brand_warning = _brand_compatible(observed_brands, canonical_brand, aliases)
    if not brand_ok:
        raise ResearchIngestionError(brand_warning or "research-observation brand does not match")
    effective_from, effective_to, warnings = _effective_dates(monitoring, source)
    if citation_placeholders:
        warnings.append(
            "citation details were not supplied for one or more monitoring questions; empty citation slices were retained"
        )
    if brand_warning:
        warnings.append(brand_warning)
    projects = _question_records(project_questions)
    records: list[dict[str, Any]] = []
    for key in sorted(monitor_keys | source_keys):
        monitor = monitoring["questions"].get(key, {})
        citation = source["questions"].get(key, {})
        project = projects.get(key)
        monitoring_available = bool(monitor)
        citation_available = bool(citation)
        variants = [*monitor.get("texts", []), *citation.get("texts", [])]
        unique_variants = list(dict.fromkeys(normal(value) for value in variants if normal(value)))
        monitor_metrics = monitor.get("research_metrics") if isinstance(monitor.get("research_metrics"), dict) else {}
        source_metrics = citation.get("research_metrics") if isinstance(citation.get("research_metrics"), dict) else {}
        records.append({
            "question_text": project["question_text"] if project else unique_variants[0],
            "question_key": key,
            "monitoring_question_texts": list(monitor.get("texts", [])),
            "source_question_texts": list(citation.get("texts", [])),
            "question_id": project.get("question_id") if project else None,
            "binding_status": "bound" if project else "unbound",
            "answer_observation_count": int(monitor.get("answer_observation_count", 0)),
            "platform_count": len(monitor.get("platforms", [])),
            "citation_detail_count": int(citation.get("citation_detail_count", 0)),
            "monitoring_available": monitoring_available,
            "citation_available": citation_available,
            "e2_ready": monitoring_available and citation_available,
            "research_metrics": {
                "brand_visibility": monitor_metrics.get("brand_visibility"),
                "citation_landscape": {
                    "citation_domain_count": source_metrics.get("citation_domain_count"),
                    "top_citation_domains": list(source_metrics.get("top_citation_domains") or []),
                    "citation_domain_concentration_top3": source_metrics.get("citation_domain_concentration_top3"),
                } if citation_available else None,
                "repeated_themes": list(monitor_metrics.get("repeated_themes") or []),
                "warnings": list(dict.fromkeys([
                    *list(monitor_metrics.get("warnings") or []),
                    *list(source_metrics.get("warnings") or []),
                ])),
            },
        })
    project_keys = set(projects)
    ready_project_keys = {item["question_key"] for item in records if item["binding_status"] == "bound" and item["e2_ready"]}
    if not ready_project_keys:
        status = "unbound"
        if project_keys & (monitor_keys | source_keys):
            status = "partial"
    elif ready_project_keys == project_keys:
        status = "ready"
    else:
        status = "partial"
    if status == "unbound":
        warnings.append("research observation questions do not exactly match the current project optimization targets")
    elif status == "partial":
        warnings.append("research observation pair covers only part of the current project optimization targets")
    return {
        "dataset_id": dataset_id,
        "content_fingerprint": content_fingerprint,
        "ingested_at": ingested_at,
        "status": status,
        "monitoring_answers_path": monitoring_member,
        "source_workbook_path": source_member,
        "monitoring_observed_from": monitoring.get("observed_from"),
        "monitoring_observed_to": monitoring.get("observed_to"),
        "source_observed_from": source.get("observed_from"),
        "source_observed_to": source.get("observed_to"),
        "effective_observed_from": effective_from,
        "effective_observed_to": effective_to,
        "declared_period_label": declared_period_label,
        "platforms": monitoring.get("platforms", []),
        "questions": records,
        "ready_question_count": sum(1 for item in records if item["e2_ready"]),
        "warnings": list(dict.fromkeys(warnings)),
    }


def empty_index() -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "artifact_type": ARTIFACT_TYPE,
        "usage_policy": USAGE_POLICY,
        "datasets": [],
    }


def load_index(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return empty_index()
    payload = _strict_json_bytes(path.read_bytes(), str(path))
    errors, _warnings = validate_index_payload(payload)
    if errors:
        raise ResearchIngestionError("invalid research-ingestion index: " + "; ".join(errors))
    return payload


def write_index(path: Path, index: dict[str, Any]) -> None:
    errors, _warnings = validate_index_payload(index)
    if errors:
        raise ResearchIngestionError("refusing to write invalid research-ingestion index: " + "; ".join(errors))
    _atomic_json(path, index)


def _fingerprint(monitoring: bytes, source: bytes) -> tuple[str, str]:
    digest = hashlib.sha256()
    digest.update(b"frontmind-question-research-ingestion-v1\0monitoring\0")
    digest.update(monitoring)
    digest.update(b"\0source\0")
    digest.update(source)
    hexdigest = digest.hexdigest()
    return f"sha256:{hexdigest}", f"observation_{hexdigest[:20]}"


def _staged_dataset_matches(destination: Path, dataset: dict[str, Any], monitoring: bytes, source: bytes) -> bool:
    if destination.is_symlink() or not destination.is_dir():
        return False
    monitor_name = PurePosixPath(dataset["monitoring_answers_path"]).name
    monitor_path = destination / monitor_name
    source_path = destination / "source_workbook.xlsx"
    if monitor_path.is_symlink() or source_path.is_symlink() or not monitor_path.is_file() or not source_path.is_file():
        return False
    return monitor_path.read_bytes() == monitoring and source_path.read_bytes() == source


def ingest_pairs(
    index: dict[str, Any] | None,
    pairs: Sequence[tuple[Path, Path]],
    staging_root: Path,
    project_questions: Any,
    canonical_brand: str,
    aliases: Sequence[str] = (),
    *,
    ingested_at: str | None = None,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Validate and stage complete pairs, returning a new index plus actions.

    Every pair is fully inventoried before the first file is copied.  Dataset
    members are content addressed and the caller persists the returned index
    with :func:`write_index` after the caller's surrounding transaction succeeds.
    """

    current = json.loads(json.dumps(index or empty_index(), ensure_ascii=False))
    errors, _warnings = validate_index_payload(current)
    if errors:
        raise ResearchIngestionError("existing research-ingestion index is invalid: " + "; ".join(errors))
    timestamp = ingested_at or datetime.now(timezone.utc).isoformat()
    existing_by_fingerprint = {item["content_fingerprint"]: item for item in current["datasets"]}
    prepared: list[tuple[Path, Path, dict[str, Any], bytes, bytes]] = []
    actions: list[dict[str, Any]] = []
    for monitoring_input, source_input in pairs:
        monitoring_path = _regular_file(Path(monitoring_input), MONITORING_SUFFIXES, "research-observation monitoring answers")
        source_path = _regular_file(Path(source_input), {".xlsx"}, "research-observation source workbook")
        if monitoring_path == source_path:
            raise ResearchIngestionError("monitoring answers and source workbook must be different files")
        monitoring_bytes = monitoring_path.read_bytes()
        source_bytes = source_path.read_bytes()
        fingerprint, dataset_id = _fingerprint(monitoring_bytes, source_bytes)
        if fingerprint in existing_by_fingerprint or any(item[2]["content_fingerprint"] == fingerprint for item in prepared):
            existing = existing_by_fingerprint.get(fingerprint)
            actions.append({"action": "duplicate_skipped", "dataset_id": (existing or {"dataset_id": dataset_id})["dataset_id"]})
            continue
        monitoring_inventory = _monitoring_inventory(monitoring_path, canonical_brand, aliases)
        source_inventory = _source_inventory(source_path)
        monitor_name = "monitoring_answers" + monitoring_path.suffix.casefold()
        monitor_member = f"research/observations/{dataset_id}/{monitor_name}"
        source_member = f"research/observations/{dataset_id}/source_workbook.xlsx"
        dataset = _build_dataset(
            monitoring_inventory,
            source_inventory,
            dataset_id=dataset_id,
            content_fingerprint=fingerprint,
            monitoring_member=monitor_member,
            source_member=source_member,
            project_questions=project_questions,
            ingested_at=timestamp,
            declared_period_label=f"{monitoring_path.stem} | {source_path.stem}",
            canonical_brand=canonical_brand,
            aliases=aliases,
        )
        prepared.append((monitoring_path, source_path, dataset, monitoring_bytes, source_bytes))

    stage = staging_root.resolve()
    stage.mkdir(parents=True, exist_ok=True)
    newly_created: list[Path] = []
    try:
        for _monitoring_path, _source_path, dataset, monitoring_bytes, source_bytes in prepared:
            destination = stage / dataset["dataset_id"]
            if destination.exists():
                if not _staged_dataset_matches(destination, dataset, monitoring_bytes, source_bytes):
                    raise ResearchIngestionError(f"staging dataset directory conflicts with submitted content: {destination.name}")
                current["datasets"].append(dataset)
                actions.append({"action": "recovered", "dataset_id": dataset["dataset_id"], "status": dataset["status"]})
                continue
            temporary = Path(tempfile.mkdtemp(prefix=f".{dataset['dataset_id']}.", dir=stage))
            try:
                monitoring_name = PurePosixPath(dataset["monitoring_answers_path"]).name
                (temporary / monitoring_name).write_bytes(monitoring_bytes)
                (temporary / "source_workbook.xlsx").write_bytes(source_bytes)
                os.replace(temporary, destination)
                newly_created.append(destination)
            except Exception:
                shutil.rmtree(temporary, ignore_errors=True)
                raise
            current["datasets"].append(dataset)
            actions.append({"action": "added", "dataset_id": dataset["dataset_id"], "status": dataset["status"]})
    except Exception:
        for path in newly_created:
            shutil.rmtree(path, ignore_errors=True)
        raise
    current["datasets"] = sorted(current["datasets"], key=lambda item: (str(item.get("ingested_at") or ""), item["dataset_id"]))
    return current, actions


def bind_project_questions(index: dict[str, Any], project_questions: Any) -> dict[str, Any]:
    projects = _question_records(project_questions)
    updated = json.loads(json.dumps(index, ensure_ascii=False))
    for dataset in updated.get("datasets", []):
        ready_projects: set[str] = set()
        seen_projects: set[str] = set()
        for item in dataset.get("questions", []):
            project = projects.get(item.get("question_key"))
            item["question_id"] = project.get("question_id") if project else None
            item["binding_status"] = "bound" if project else "unbound"
            if project:
                item["question_text"] = project["question_text"]
                seen_projects.add(item["question_key"])
                if item.get("e2_ready"):
                    ready_projects.add(item["question_key"])
        project_keys = set(projects)
        if ready_projects and ready_projects == project_keys:
            dataset["status"] = "ready"
        elif ready_projects or seen_projects:
            dataset["status"] = "partial"
        else:
            dataset["status"] = "unbound"
        warnings = [item for item in dataset.get("warnings", []) if not item.startswith("research observation questions do not") and not item.startswith("research observation pair covers only")]
        if dataset["status"] == "unbound":
            warnings.append("research observation questions do not exactly match the current project optimization targets")
        elif dataset["status"] == "partial":
            warnings.append("research observation pair covers only part of the current project optimization targets")
        dataset["warnings"] = list(dict.fromkeys(warnings))
    return updated


def remove_dataset(index: dict[str, Any], dataset_id: str) -> tuple[dict[str, Any], dict[str, Any]]:
    target = normal(dataset_id)
    updated = json.loads(json.dumps(index, ensure_ascii=False))
    removed = [item for item in updated.get("datasets", []) if item.get("dataset_id") == target]
    if len(removed) != 1:
        raise ResearchIngestionError(f"unknown research-observation dataset_id: {target}")
    updated["datasets"] = [item for item in updated["datasets"] if item.get("dataset_id") != target]
    return updated, removed[0]


def select_latest_dataset(index: dict[str, Any], question: Any) -> dict[str, Any] | None:
    key = question_key(question)
    candidates: list[dict[str, Any]] = []
    for dataset in index.get("datasets", []):
        if any(item.get("question_key") == key and item.get("e2_ready") is True for item in dataset.get("questions", [])):
            candidates.append(dataset)
    # Stable sorts express the mixed directions exactly: newest observed
    # period, then newest ingestion, then ascending content-addressed ID.
    candidates.sort(key=lambda item: str(item.get("dataset_id") or ""))
    candidates.sort(key=lambda item: str(item.get("ingested_at") or ""), reverse=True)
    candidates.sort(key=lambda item: str(item.get("effective_observed_to") or ""), reverse=True)
    return json.loads(json.dumps(candidates[0], ensure_ascii=False)) if candidates else None


def summarize_index(index: dict[str, Any], project_questions: Any) -> dict[str, Any]:
    projects = _question_records(project_questions)
    datasets: list[dict[str, Any]] = []
    for dataset in index.get("datasets", []):
        questions = dataset.get("questions", [])
        datasets.append({
            "dataset_id": dataset.get("dataset_id"),
            "status": dataset.get("status"),
            "declared_period_label": dataset.get("declared_period_label"),
            "effective_observed_from": dataset.get("effective_observed_from"),
            "effective_observed_to": dataset.get("effective_observed_to"),
            "platform_count": len(dataset.get("platforms", [])),
            "answer_observation_count": sum(int(item.get("answer_observation_count") or 0) for item in questions),
            "citation_detail_count": sum(int(item.get("citation_detail_count") or 0) for item in questions),
            "common_question_count": sum(1 for item in questions if item.get("e2_ready") is True),
            "bound_question_count": sum(1 for item in questions if item.get("binding_status") == "bound"),
            "unbound_question_count": sum(1 for item in questions if item.get("binding_status") == "unbound"),
            "e2_ready_project_questions": [item["question_text"] for item in questions if item.get("binding_status") == "bound" and item.get("e2_ready")],
            "warnings": dataset.get("warnings", []),
        })
    coverage: list[dict[str, Any]] = []
    for key, project in sorted(projects.items(), key=lambda item: item[1]["source_order"]):
        monitoring = False
        citations = False
        for dataset in index.get("datasets", []):
            for item in dataset.get("questions", []):
                if item.get("question_key") == key:
                    monitoring = monitoring or item.get("monitoring_available") is True
                    citations = citations or item.get("citation_available") is True
        latest = select_latest_dataset(index, project["question_text"])
        coverage.append({
            "question_id": project.get("question_id"),
            "question_text": project["question_text"],
            "coverage": "both" if monitoring and citations else "monitoring_only" if monitoring else "citation_only" if citations else "missing",
            "latest_ready_dataset_id": latest.get("dataset_id") if latest else None,
            "latest_observed_to": latest.get("effective_observed_to") if latest else None,
        })
    return {
        "usage_policy": USAGE_POLICY,
        "dataset_count": len(datasets),
        "datasets": datasets,
        "project_question_coverage": coverage,
    }


def _controlled_research_warnings(dataset: dict[str, Any]) -> list[str]:
    """Return Pack-safe warnings without copying filenames or raw values."""

    warnings: list[str] = []
    if not dataset.get("effective_observed_from") or not dataset.get("effective_observed_to"):
        warnings.append("actual_observed_period_is_incomplete")
    if dataset.get("status") == "partial":
        warnings.append("partial_project_question_overlap")
    elif dataset.get("status") == "unbound":
        warnings.append("no_project_question_overlap")
    if any(item.get("e2_ready") is not True for item in dataset.get("questions", [])):
        warnings.append("some_questions_are_present_on_only_one_side")
    return warnings


def validate_index_payload(index: Any, available_members: set[str] | None = None) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not isinstance(index, dict):
        return ["research-ingestion index root must be an object"], warnings
    if index.get("schema_version") != SCHEMA_VERSION:
        errors.append(f"unsupported research-ingestion index schema_version: {index.get('schema_version')!r}")
    if index.get("artifact_type") != ARTIFACT_TYPE:
        errors.append(f"unexpected research-ingestion artifact_type: {index.get('artifact_type')!r}")
    if index.get("usage_policy") != USAGE_POLICY:
        errors.append("research-ingestion usage_policy must prohibit factual-evidence use")
    datasets = index.get("datasets")
    if not isinstance(datasets, list):
        errors.append("research-ingestion datasets must be an array")
        return errors, warnings
    ids: set[str] = set()
    fingerprints: set[str] = set()
    paths: set[str] = set()
    for position, dataset in enumerate(datasets):
        label = f"datasets[{position}]"
        if not isinstance(dataset, dict):
            errors.append(f"{label} must be an object")
            continue
        dataset_id = dataset.get("dataset_id")
        if not isinstance(dataset_id, str) or not re.fullmatch(r"observation_[0-9a-f]{20}", dataset_id):
            errors.append(f"{label}.dataset_id is invalid")
        elif dataset_id in ids:
            errors.append(f"duplicate research-observation dataset_id: {dataset_id}")
        else:
            ids.add(dataset_id)
        fingerprint = dataset.get("content_fingerprint")
        if not isinstance(fingerprint, str) or not re.fullmatch(r"sha256:[0-9a-f]{64}", fingerprint):
            errors.append(f"{label}.content_fingerprint is invalid")
        elif fingerprint in fingerprints:
            errors.append(f"duplicate research-observation content_fingerprint: {fingerprint}")
        else:
            fingerprints.add(fingerprint)
        ingested_at = dataset.get("ingested_at")
        try:
            parsed_ingested = datetime.fromisoformat(str(ingested_at).replace("Z", "+00:00"))
            if parsed_ingested.tzinfo is None or parsed_ingested.utcoffset() is None:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"{label}.ingested_at must be an ISO date-time with timezone")
        if dataset.get("status") not in {"ready", "partial", "unbound"}:
            errors.append(f"{label}.status is invalid")
        platforms = dataset.get("platforms")
        if not isinstance(platforms, list) or any(not isinstance(item, str) or not item for item in platforms) or len(platforms) != len(set(platforms)):
            errors.append(f"{label}.platforms must contain unique non-empty strings")
        dataset_warnings = dataset.get("warnings")
        if not isinstance(dataset_warnings, list) or any(not isinstance(item, str) or not item for item in dataset_warnings):
            errors.append(f"{label}.warnings must contain strings")
        for role in ("monitoring_answers_path", "source_workbook_path"):
            value = dataset.get(role)
            prefix = f"research/observations/{dataset_id}/" if isinstance(dataset_id, str) else "research/observations/"
            if not safe_relative(value) or not str(value).startswith(prefix):
                errors.append(f"{label}.{role} is unsafe or outside its dataset directory")
                continue
            if value in paths:
                errors.append(f"research-observation member is assigned more than once: {value}")
            paths.add(value)
            if available_members is not None and value not in available_members:
                errors.append(f"declared research-observation member is absent: {value}")
        if dataset.get("monitoring_answers_path") == dataset.get("source_workbook_path"):
            errors.append(f"{label} assigns both roles to the same member")
        questions = dataset.get("questions")
        if not isinstance(questions, list) or not questions:
            errors.append(f"{label}.questions must be a non-empty array")
            continue
        keys: set[str] = set()
        ready_count = 0
        for question_position, question in enumerate(questions):
            question_label = f"{label}.questions[{question_position}]"
            if not isinstance(question, dict):
                errors.append(f"{question_label} must be an object")
                continue
            key = question.get("question_key")
            if not isinstance(key, str) or not key or question_key(question.get("question_text")) != key:
                errors.append(f"{question_label}.question_key does not match question_text")
            elif key in keys:
                errors.append(f"{label} contains duplicate question_key: {key}")
            keys.add(str(key))
            expected_ready = question.get("monitoring_available") is True and question.get("citation_available") is True
            if question.get("e2_ready") is not expected_ready:
                errors.append(f"{question_label}.e2_ready is inconsistent with side availability")
            if expected_ready:
                ready_count += 1
            if question.get("binding_status") not in {"bound", "unbound"}:
                errors.append(f"{question_label}.binding_status is invalid")
            for variants_field in ("monitoring_question_texts", "source_question_texts"):
                variants = question.get(variants_field)
                if not isinstance(variants, list) or any(not isinstance(item, str) or not item for item in variants):
                    errors.append(f"{question_label}.{variants_field} must contain strings")
            for count_field in ("answer_observation_count", "platform_count", "citation_detail_count"):
                count = question.get(count_field)
                if isinstance(count, bool) or not isinstance(count, int) or count < 0:
                    errors.append(f"{question_label}.{count_field} must be a non-negative integer")
        if dataset.get("ready_question_count") != ready_count:
            errors.append(f"{label}.ready_question_count does not match question records")
        if ready_count == 0:
            errors.append(f"{label} has no question available on both sides")
        for date_field in (
            "monitoring_observed_from", "monitoring_observed_to", "source_observed_from", "source_observed_to",
            "effective_observed_from", "effective_observed_to",
        ):
            value = dataset.get(date_field)
            if value is not None:
                try:
                    date.fromisoformat(str(value))
                except ValueError:
                    errors.append(f"{label}.{date_field} is not an ISO date or null")
        if dataset.get("effective_observed_from") and dataset.get("effective_observed_to") and dataset["effective_observed_from"] > dataset["effective_observed_to"]:
            errors.append(f"{label} effective observed range is reversed")
        if not dataset.get("effective_observed_to"):
            warnings.append(f"{dataset_id or label} has no effective observed_to; ingested_at is the fallback")
    return errors, warnings




def staging_member_path(staging_root: Path, member: str) -> Path:
    """Map one historical monitoring/citation member into a private staging root."""

    prefix = "research/observations/"
    if not safe_relative(member) or not member.startswith(prefix) or member == INDEX_MEMBER:
        raise ResearchIngestionError(f"not a canonical staged research member: {member!r}")
    relative = PurePosixPath(member).relative_to(prefix)
    destination = staging_root.resolve().joinpath(*relative.parts)
    if staging_root.resolve() not in destination.parents:
        raise ResearchIngestionError(f"research member escapes staging root: {member}")
    return destination


__all__ = [
    "ARTIFACT_TYPE", "BRAND_ALIASES", "DATE_ALIASES", "DETAIL_SHEET_ALIASES",
    "INDEX_MEMBER", "MONITORING_SUFFIXES", "PLATFORM_ALIASES", "QUESTION_ALIASES",
    "ResearchIngestionError", "URL_ALIASES", "_DirectoryPackageReader",
    "_ZipPackageReader", "_answer_values", "_brand_compatible", "_date_values",
    "_effective_dates", "_http_url", "_monitoring_inventory", "_regular_file",
    "_rows_from_csv", "_rows_from_xlsx", "_source_inventory", "_strict_json_bytes",
    "empty_index", "ingest_pairs", "field_key", "normal", "safe_relative", "staging_member_path",
    "validate_index_payload",
]
