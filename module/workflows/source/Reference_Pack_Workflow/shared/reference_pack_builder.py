#!/usr/bin/env python3
"""Safe construction and versioned maintenance of Reference Pack 4.1.

The builder owns intake, canonical storage, portable archives and immutable
version commits. Positioning judgment remains a provider action in the
controller, but its confirmed research and strategy artifacts are committed
into the same Reference Pack instead of a second Foundation package.
"""
from __future__ import annotations

import csv
import hashlib
import html
import json
import os
import re
import shutil
import stat
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Mapping, Sequence

from shared import research_ingestion
from shared.question_research import (
    QuestionResearchError,
    build_from_research_ingestion_index,
    load_question_research_index,
)
from shared.reference_pack import BRAND_MARKET_MEMBERS, P0_MEMBERS, STRATEGY_MEMBERS, current_root
from shared.scripts.validate_execution_reference_pack import DirectoryView, validate as validate_pack
from Reference_Pack_Workflow.shared.intake_security import (
    IntakeBudget,
    IntakeSecurityError,
    MAX_COMPRESSION_RATIO,
    MAX_EXTRACTED_TEXT_CHARS_PER_JOB,
    MAX_INPUT_BYTES,
    MAX_INPUT_FILES,
    MAX_PDF_PAGES,
    MAX_PPTX_SLIDES,
    MAX_SPREADSHEET_NONEMPTY_CELLS,
    MAX_SPREADSHEET_SHEETS,
    MAX_TOTAL_BYTES,
    enforce_text_budget,
    freeze_regular_file,
    path_has_symlink_component,
    validate_material,
    validate_zip_infos,
)


BUILDER_VERSION = "frontmind-reference-pack-builder/4.11"
TEXT_SUFFIXES = {
    ".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".html", ".htm",
}
DOCUMENT_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx", ".xls"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".svg"}
ACCEPTED_SUFFIXES = TEXT_SUFFIXES | DOCUMENT_SUFFIXES | IMAGE_SUFFIXES
ARCHIVE_SUFFIXES = {".zip"}
SYSTEM_NAMES = {".DS_Store", "Thumbs.db"}
ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


class ReferencePackBuildError(ValueError):
    """A supplied knowledge-base item cannot safely become a Pack fact asset."""


@dataclass
class MaterialCandidate:
    path: Path
    source_path: str
    classification: str
    report_index: int
    sha256: str = ""
    aliases: list[str] = field(default_factory=list)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def _atomic_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(value if value.endswith("\n") else value + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def _atomic_bytes(path: Path, value: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(value)
            handle.flush()
            os.fsync(handle.fileno())
        os.chmod(temporary, 0o644)
        os.replace(temporary, path)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def _normal(value: Any) -> str:
    return re.sub(
        r"[ \t\u3000]+", " ",
        str(value or "").replace("\r\n", "\n").replace("\r", "\n"),
    ).strip()


def _safe_name(name: str, fallback: str) -> str:
    stem = re.sub(r"[^0-9A-Za-z_\-\u3400-\u9fff.]", "_", name).strip("._")
    return stem[:120] or fallback


def _safe_relative(value: str) -> bool:
    if not isinstance(value, str) or not value or "\\" in value or "\x00" in value:
        return False
    path = PurePosixPath(value)
    return not path.is_absolute() and ".." not in path.parts and path.as_posix() == value


def _system_garbage(value: str) -> bool:
    parts = PurePosixPath(value).parts
    return "__MACOSX" in parts or bool(parts and parts[-1] in SYSTEM_NAMES)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while block := handle.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def _freeze_regular_file(
    path: Path, staging: Path, source_path: str, *, budget: IntakeBudget,
) -> tuple[Path, int, str]:
    """Map the central intake error type to the Pack builder's public error."""

    try:
        return freeze_regular_file(
            path, staging, label=source_path, budget=budget,
            budget_label="knowledge_base_job",
            enforce_size_limit=False,
        )
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc


def _item(
    reports: list[dict[str, Any]], source: str, status: str, classification: str,
    *, reason: str | None = None, sha256: str | None = None,
    aliases: Sequence[str] = (),
) -> int:
    value: dict[str, Any] = {
        "source_path": source,
        "status": status,
        "classification": classification,
    }
    if reason:
        value["reason"] = reason
    if sha256:
        value["sha256"] = f"sha256:{sha256}"
    if aliases:
        value["aliases"] = list(dict.fromkeys(aliases))
    reports.append(value)
    return len(reports) - 1


def _is_canonical_pack_directory(path: Path) -> bool:
    return bool(
        (path / "reference_pack.json").is_file()
        and (path / "materials/index.json").is_file()
        and (path / "registries/source_registry.json").is_file()
    )


def _zip_prefix_for(names: Sequence[str], member: str) -> str | None:
    if member in names:
        return ""
    matches = [name[:-len(member)] for name in names if name.endswith("/" + member)]
    return matches[0] if len(matches) == 1 else None


def _decode_manifest(data: bytes) -> dict[str, Any] | None:
    try:
        value = json.loads(data.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return None
    return value if isinstance(value, dict) else None


def _classify_zip(archive: zipfile.ZipFile, infos: Sequence[zipfile.ZipInfo]) -> str:
    names = [info.filename for info in infos if not info.is_dir()]
    lowered = [name.casefold() for name in names]
    roots = {PurePosixPath(name).parts[0].casefold() for name in names if PurePosixPath(name).parts}
    if (
        any(root.startswith("frontmind_content_workflow") for root in roots)
        or any(name.endswith("/scripts/frontmind") or name == "scripts/frontmind" for name in lowered)
        or (
            any("reference_pack_workflow/" in name for name in lowered)
            and any("execution_workflow/" in name for name in lowered)
        )
    ):
        return "workflow_release_zip"
    pack_prefix = _zip_prefix_for(names, "reference_pack.json")
    if pack_prefix is not None and all(
        pack_prefix + member in names
        for member in ("materials/index.json", "registries/source_registry.json")
    ):
        return "reference_pack_zip"
    manifests = [
        info for info in infos
        if not info.is_dir() and PurePosixPath(info.filename).name in {
            "00_package_manifest.json", "manifest.json", "frontmind.kb-working-set.json",
        }
    ]
    for info in manifests:
        if info.file_size > min(MAX_INPUT_BYTES, 16 * 1024 * 1024):
            continue
        manifest = _decode_manifest(archive.read(info))
        if manifest is None:
            continue
        identity = " ".join(str(manifest.get(key) or "") for key in ("kind", "profile", "schema", "format"))
        if "frontmind.kb-working-set" in identity:
            return "working_set_zip"
        if (
            manifest.get("kind") == "frontmind.dashboard-owned-knowledge-package"
            or "knowledge-package" in identity
            or "knowledge_base" in identity
        ):
            return "canonical_knowledge_base_zip"
    usable = [name for name in names if PurePosixPath(name).suffix.casefold() in ACCEPTED_SUFFIXES]
    if usable and all("working-set" in PurePosixPath(name).parts for name in usable):
        return "working_set_zip"
    return "material_zip"


def _validate_zip_metadata(
    infos: Sequence[zipfile.ZipInfo], *, budget: IntakeBudget,
) -> None:
    """Apply structural ZIP checks without outer/member/total byte ceilings.

    v4.11 keeps the 5,000-member and compression-ratio protections. Byte size
    is left to streaming I/O and the host filesystem, so a large but otherwise
    valid Reference Pack or knowledge-material archive is not rejected merely
    because it crossed the old 128/512 MiB thresholds.
    """
    try:
        validate_zip_infos(
            infos, label="knowledge_base_archive", budget=None,
            enforce_size_limits=False,
        )
        next_count = budget.file_count + len(infos)
        if next_count > MAX_INPUT_FILES:
            raise IntakeSecurityError("knowledge_base_archive_too_many_members")
        budget.file_count = next_count
        files = [item for item in infos if not item.is_dir()]
        budget.uncompressed_bytes += sum(item.file_size for item in files)
        budget.compressed_bytes += sum(item.compress_size for item in files)
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc


def _stream_copy(source: BinaryIO, destination: Path, expected: int) -> str:
    destination.parent.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    total = 0
    with destination.open("wb") as target:
        while block := source.read(1024 * 1024):
            total += len(block)
            if total > expected:
                raise ReferencePackBuildError("knowledge_base_archive_member_size_mismatch")
            digest.update(block)
            target.write(block)
    if total != expected:
        raise ReferencePackBuildError("knowledge_base_archive_member_size_mismatch")
    return digest.hexdigest()


def _collect_zip(
    path: Path, staging: Path, reports: list[dict[str, Any]], candidates: list[MaterialCandidate],
    *, budget: IntakeBudget, source_label: str | None = None,
) -> None:
    label = source_label or path.name
    local_reports: list[dict[str, Any]] = []
    local_candidates: list[MaterialCandidate] = []
    created_paths: list[Path] = []
    archive_budget = IntakeBudget(
        file_count=budget.file_count,
        uncompressed_bytes=budget.uncompressed_bytes,
        compressed_bytes=budget.compressed_bytes,
    )
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            _validate_zip_metadata(infos, budget=archive_budget)
            classification = _classify_zip(archive, infos)
            if classification == "workflow_release_zip":
                _item(reports, label, "rejected", classification, reason="workflow_release_zip_not_knowledge_base")
                return
            if classification == "reference_pack_zip":
                _item(
                    reports, label, "rejected", classification,
                    reason="canonical_reference_pack_requires_existing_pack_route",
                )
                return
            nested = next(
                (info.filename for info in infos if not info.is_dir() and PurePosixPath(info.filename).suffix.casefold() == ".zip"),
                None,
            )
            if nested:
                raise ReferencePackBuildError(
                    f"knowledge_base_archive_nested_archive_forbidden: {nested}"
                )
            archive_index = _item(local_reports, label, "accepted", classification)
            del archive_index
            usable = 0
            for info in sorted(infos, key=lambda value: value.filename):
                if info.is_dir():
                    continue
                identity = f"{label}!/{info.filename}"
                if _system_garbage(info.filename):
                    _item(local_reports, identity, "skipped", "archive_member", reason="system_metadata")
                    continue
                suffix = PurePosixPath(info.filename).suffix.casefold()
                if suffix not in ACCEPTED_SUFFIXES:
                    _item(local_reports, identity, "skipped", "archive_member", reason="unsupported_file_type")
                    continue
                target = staging / (
                    f"zip_{len(candidates) + len(local_candidates):05d}_"
                    f"{_safe_name(PurePosixPath(info.filename).name, 'material')}"
                )
                created_paths.append(target)
                try:
                    with archive.open(info, "r") as source:
                        digest = _stream_copy(source, target, info.file_size)
                except (OSError, RuntimeError, zipfile.BadZipFile, ReferencePackBuildError) as exc:
                    raise ReferencePackBuildError(
                        f"knowledge_base_archive_member_unreadable: {info.filename}: {exc}"
                    ) from exc
                report_index = _item(
                    local_reports, identity, "accepted", "archive_member",
                    sha256=digest, aliases=[identity],
                )
                local_candidates.append(MaterialCandidate(
                    target, identity, classification, report_index,
                    sha256=digest, aliases=[identity],
                ))
                usable += 1
            if not usable:
                for item in local_reports:
                    if item.get("source_path") == label and item.get("classification") == classification:
                        item["status"] = "rejected"
                        item["reason"] = "knowledge_base_archive_has_no_usable_materials"
                        break
            # Commit metadata, candidates and the cumulative archive budget as
            # one unit only after every member was streamed and CRC-checked.
            report_offset = len(reports)
            for candidate in local_candidates:
                candidate.report_index += report_offset
            reports.extend(local_reports)
            candidates.extend(local_candidates)
            budget.file_count = archive_budget.file_count
            budget.uncompressed_bytes = archive_budget.uncompressed_bytes
            budget.compressed_bytes = archive_budget.compressed_bytes
    except ReferencePackBuildError as exc:
        for created in created_paths:
            created.unlink(missing_ok=True)
        _item(reports, label, "rejected", "zip", reason=str(exc))
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        for created in created_paths:
            created.unlink(missing_ok=True)
        _item(reports, label, "rejected", "zip", reason=f"knowledge_base_archive_unreadable: {exc}")


def _collect_file(
    path: Path, source_path: str, classification: str,
    reports: list[dict[str, Any]], candidates: list[MaterialCandidate],
    *, budget: IntakeBudget, staging: Path,
) -> None:
    if _system_garbage(source_path):
        _item(reports, source_path, "skipped", classification, reason="system_metadata")
        return
    suffix = path.suffix.casefold()
    if suffix == ".zip":
        _item(reports, source_path, "rejected", classification, reason="nested_archive_not_expanded")
        return
    if path.name == "reference_pack.json":
        _item(
            reports, source_path, "rejected", classification,
            reason="canonical_reference_pack_manifest_not_knowledge_material",
        )
        return
    if suffix not in ACCEPTED_SUFFIXES:
        _item(reports, source_path, "rejected", classification, reason="knowledge_base_unsupported_file_type")
        return
    budget_before = (
        budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes,
    )
    frozen: Path | None = None
    try:
        frozen, size, digest = _freeze_regular_file(
            path, staging, source_path, budget=budget,
        )
    except (IntakeSecurityError, ReferencePackBuildError) as exc:
        if frozen is not None:
            frozen.unlink(missing_ok=True)
        budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes = budget_before
        _item(reports, source_path, "rejected", classification, reason=str(exc))
        return
    report_index = _item(
        reports, source_path, "accepted", classification,
        sha256=digest, aliases=[source_path],
    )
    candidates.append(MaterialCandidate(
        frozen, source_path, classification, report_index,
        sha256=digest, aliases=[source_path],
    ))


def _collect_inputs(
    inputs: Sequence[Path], staging: Path, *, budget: IntakeBudget,
) -> tuple[list[MaterialCandidate], list[dict[str, Any]]]:
    candidates: list[MaterialCandidate] = []
    reports: list[dict[str, Any]] = []
    encountered = 0
    budget_at_start = (
        budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes,
    )
    global_fatal_reason: str | None = None

    def budget_or_count_failure(reason: str) -> bool:
        return reason.startswith((
            "knowledge_base_job_", "knowledge_base_archive_too_many",
            "knowledge_base_archive_total_size", "knowledge_base_archive_compression",
            "knowledge_base_has_too_many_files",
        ))

    for raw in inputs:
        requested = Path(raw).expanduser()
        label = requested.name or str(requested)
        if path_has_symlink_component(requested) or requested.is_symlink() or not requested.exists():
            _item(reports, label, "rejected", "input", reason="knowledge_base_input_missing_or_symlink")
            continue
        source = requested.resolve()
        if source.is_file():
            report_start = len(reports)
            if source.suffix.casefold() == ".zip":
                try:
                    frozen_archive, _, _ = freeze_regular_file(
                        source, staging, label=label, budget=None,
                        enforce_size_limit=False,
                    )
                except (IntakeSecurityError, ReferencePackBuildError) as exc:
                    _item(reports, label, "rejected", "zip", reason=str(exc))
                else:
                    try:
                        _collect_zip(
                            frozen_archive, staging, reports, candidates,
                            budget=budget, source_label=label,
                        )
                    finally:
                        frozen_archive.unlink(missing_ok=True)
            else:
                _collect_file(
                    source, label, "single_file", reports, candidates,
                    budget=budget, staging=staging,
                )
            encountered += 1
            new_reasons = [
                str(item.get("reason") or "") for item in reports[report_start:]
                if item.get("status") == "rejected"
            ]
            global_fatal_reason = next(
                (reason for reason in new_reasons if budget_or_count_failure(reason)), None,
            )
            if global_fatal_reason:
                break
            continue
        if not source.is_dir():
            _item(reports, label, "rejected", "input", reason="knowledge_base_item_not_regular")
            continue
        if _is_canonical_pack_directory(source):
            _item(
                reports, label, "rejected", "reference_pack_directory",
                reason="canonical_reference_pack_requires_existing_pack_route",
            )
            continue
        directory_index = _item(reports, label, "accepted", "directory")
        before = len(candidates)
        report_before_members = len(reports)
        budget_before = (
            budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes,
        )
        fatal_reason: str | None = None
        for item in sorted(source.rglob("*"), key=lambda value: value.relative_to(source).as_posix()):
            encountered += 1
            if encountered > MAX_INPUT_FILES:
                fatal_reason = "knowledge_base_has_too_many_files"
                break
            relative = item.relative_to(source).as_posix()
            identity = f"{label}/{relative}"
            if item.is_symlink():
                fatal_reason = f"knowledge_base_directory_symlink_forbidden: {relative}"
                break
            elif item.is_file():
                if item.suffix.casefold() == ".zip":
                    _item(reports, identity, "rejected", "directory_member", reason="nested_archive_not_expanded")
                else:
                    report_count = len(reports)
                    _collect_file(
                        item, identity, "directory_member", reports, candidates,
                        budget=budget, staging=staging,
                    )
                    if len(reports) > report_count:
                        reason = str(reports[-1].get("reason") or "")
                        if reports[-1].get("status") == "rejected" and reason.startswith((
                            "knowledge_base_job_",
                            "knowledge_base_input_changed",
                            "knowledge_base_item_unreadable",
                            "knowledge_base_item_not_regular",
                            "knowledge_base_frozen_copy_hash_mismatch",
                        )):
                            fatal_reason = reason
                            break
            elif not item.is_dir():
                fatal_reason = f"knowledge_base_directory_special_file_forbidden: {relative}"
                break
        if fatal_reason:
            for candidate in candidates[before:]:
                candidate.path.unlink(missing_ok=True)
            del candidates[before:]
            del reports[report_before_members:]
            budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes = budget_before
            reports[directory_index]["status"] = "rejected"
            reports[directory_index]["reason"] = fatal_reason
            if budget_or_count_failure(fatal_reason):
                global_fatal_reason = fatal_reason
                break
        if len(candidates) == before and reports[directory_index]["status"] == "accepted":
            reports[directory_index]["status"] = "rejected"
            reports[directory_index]["reason"] = "knowledge_base_directory_has_no_usable_materials"
    if encountered > MAX_INPUT_FILES or global_fatal_reason:
        for candidate in candidates:
            candidate.path.unlink(missing_ok=True)
        candidates.clear()
        budget.file_count, budget.uncompressed_bytes, budget.compressed_bytes = budget_at_start
    return candidates, reports


def is_knowledge_base_input(path: Path) -> bool:
    """Shallow route classifier; it deliberately does not open ZIPs pre-selection."""
    try:
        requested = Path(path).expanduser()
        if requested.is_symlink() or not requested.exists():
            return False
        if requested.is_file():
            return requested.suffix.casefold() in ACCEPTED_SUFFIXES | ARCHIVE_SUFFIXES
        return requested.is_dir() and not _is_canonical_pack_directory(requested.resolve())
    except OSError:
        return False


def _repair_markdown(text: str) -> tuple[str, list[str]]:
    repairs: list[str] = []
    value = text.replace("\r\n", "\n").replace("\r", "\n").replace("\ufeff", "")
    if value != text:
        repairs.append("normalized_line_endings")
    repaired = re.sub(r"(?m)^(#{1,6})([^\s#])", r"\1 \2", value)
    if repaired != value:
        repairs.append("normalized_atx_heading_spacing")
    value = repaired
    if len(re.findall(r"(?m)^\s*```", value)) % 2:
        value = value.rstrip() + "\n```\n"
        repairs.append("closed_unmatched_code_fence")
    return value, repairs


def _decode_text(data: bytes, name: str) -> str:
    encodings = ["utf-8-sig"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.insert(0, "utf-16")
    encodings.append("gb18030")
    for encoding in encodings:
        try:
            value = data.decode(encoding)
            if "\x00" not in value:
                return value
        except UnicodeDecodeError:
            continue
    raise ReferencePackBuildError(f"knowledge_base_text_encoding_unsupported: {name}")


def _text_from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        if bool(getattr(reader, "is_encrypted", False)):
            raise ReferencePackBuildError(f"knowledge_base_pdf_encrypted: {path.name}")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ReferencePackBuildError(f"knowledge_base_pdf_page_limit_exceeded: {path.name}")
        value = "\n\n".join(
            f"# PDF page {number}\n{page.extract_text() or ''}"
            for number, page in enumerate(reader.pages, 1)
        )
        return enforce_text_budget(value, label=path.name)
    except ReferencePackBuildError:
        raise
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc
    except Exception as exc:
        raise ReferencePackBuildError(f"knowledge_base_pdf_unreadable: {path.name}") from exc


def _text_from_docx(path: Path) -> str:
    try:
        from docx import Document
        document = Document(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        lines.extend(f"[图片引用 {number}]" for number, _ in enumerate(document.inline_shapes, 1))
        return enforce_text_budget("\n".join(lines), label=path.name)
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc
    except Exception as exc:
        raise ReferencePackBuildError(f"knowledge_base_docx_unreadable: {path.name}") from exc


def _text_from_sheet(path: Path) -> str:
    suffix = path.suffix.casefold()
    try:
        lines: list[str] = []
        if suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(str(path), on_demand=True)
            try:
                if workbook.nsheets > MAX_SPREADSHEET_SHEETS:
                    raise ReferencePackBuildError(
                        f"knowledge_base_spreadsheet_sheet_limit_exceeded: {path.name}"
                    )
                nonempty_cells = 0
                for sheet in workbook.sheets():
                    lines.append(f"# Sheet: {sheet.name}")
                    for row_number in range(sheet.nrows):
                        cells = [_normal(value) for value in sheet.row_values(row_number) if _normal(value)]
                        nonempty_cells += len(cells)
                        if nonempty_cells > MAX_SPREADSHEET_NONEMPTY_CELLS:
                            raise ReferencePackBuildError(
                                f"knowledge_base_spreadsheet_cell_limit_exceeded: {path.name}"
                            )
                        if cells:
                            lines.append(" | ".join(cells))
            finally:
                workbook.release_resources()
            return enforce_text_budget("\n".join(lines), label=path.name)
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if len(workbook.worksheets) > MAX_SPREADSHEET_SHEETS:
                raise ReferencePackBuildError(
                    f"knowledge_base_spreadsheet_sheet_limit_exceeded: {path.name}"
                )
            nonempty_cells = 0
            for sheet in workbook.worksheets:
                lines.append(f"# Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = [_normal(value) for value in row if _normal(value)]
                    nonempty_cells += len(cells)
                    if nonempty_cells > MAX_SPREADSHEET_NONEMPTY_CELLS:
                        raise ReferencePackBuildError(
                            f"knowledge_base_spreadsheet_cell_limit_exceeded: {path.name}"
                        )
                    if cells:
                        lines.append(" | ".join(cells))
            return enforce_text_budget("\n".join(lines), label=path.name)
        finally:
            workbook.close()
    except ReferencePackBuildError:
        raise
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc
    except Exception as exc:
        raise ReferencePackBuildError(f"knowledge_base_spreadsheet_unreadable: {path.name}") from exc


def _slide_number(name: str) -> int:
    matched = re.search(r"slide(\d+)\.xml$", name)
    return int(matched.group(1)) if matched else 0


def _text_from_pptx(path: Path) -> str:
    try:
        with zipfile.ZipFile(path) as archive:
            names = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
            if not names:
                raise ValueError("no slide XML")
            if len(names) > MAX_PPTX_SLIDES:
                raise ReferencePackBuildError(
                    f"knowledge_base_presentation_slide_limit_exceeded: {path.name}"
                )
            lines: list[str] = []
            for name in sorted(names, key=_slide_number):
                lines.append(f"# Slide {_slide_number(name)}")
                xml = archive.read(name).decode("utf-8", errors="replace")
                lines.extend(html.unescape(item) for item in re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))
            return enforce_text_budget("\n".join(lines), label=path.name)
    except ReferencePackBuildError:
        raise
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc
    except Exception as exc:
        raise ReferencePackBuildError(f"knowledge_base_presentation_unreadable: {path.name}") from exc


def _read_textual_material(
    path: Path, *, budget: IntakeBudget | None = None,
) -> tuple[str, list[str], str]:
    suffix = path.suffix.casefold()
    try:
        validate_material(path, budget=budget)
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc
    if suffix in IMAGE_SUFFIXES:
        return "", [], "image_only"
    if suffix == ".pdf":
        value = _text_from_pdf(path)
        return value, [], "usable" if _text_blocks(value) else "needs_ocr"
    if suffix == ".docx":
        return _text_from_docx(path), [], "usable"
    if suffix in {".xlsx", ".xls"}:
        return _text_from_sheet(path), [], "usable"
    if suffix == ".pptx":
        return _text_from_pptx(path), [], "usable"
    try:
        raw = _decode_text(path.read_bytes(), path.name)
    except OSError as exc:
        raise ReferencePackBuildError(f"knowledge_base_text_unreadable: {path.name}") from exc
    if suffix in {".md", ".markdown"}:
        raw, repairs = _repair_markdown(raw)
        try:
            return enforce_text_budget(raw, label=path.name), repairs, "usable"
        except IntakeSecurityError as exc:
            raise ReferencePackBuildError(str(exc)) from exc
    if suffix in {".html", ".htm"}:
        raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
        raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        raw = html.unescape(raw)
    elif suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            rows = csv.reader(raw.splitlines(), delimiter=delimiter, strict=True)
            raw = "\n".join(
                " | ".join(_normal(cell) for cell in row if _normal(cell))
                for row in rows
            )
        except csv.Error as exc:
            raise ReferencePackBuildError(f"knowledge_base_delimited_text_invalid: {path.name}") from exc
    elif suffix == ".json":
        try:
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2, sort_keys=True)
        except json.JSONDecodeError as exc:
            raise ReferencePackBuildError(f"knowledge_base_json_invalid: {path.name}") from exc
    try:
        return enforce_text_budget(raw, label=path.name), [], "usable"
    except IntakeSecurityError as exc:
        raise ReferencePackBuildError(str(exc)) from exc


def extract_safe_material_text(
    path: Path, *, budget: IntakeBudget | None = None,
) -> dict[str, Any]:
    """Extract one material through the canonical v4.11 intake security gates.

    This narrow API is shared by Pack construction and ordinary user-material
    intake.  Callers therefore cannot accidentally bypass MIME, OOXML, page,
    slide, sheet, cell, image, decoder-bomb, or extracted-text limits by
    implementing a second parser path.
    """

    requested = Path(path).expanduser()
    if requested.is_symlink() or not requested.is_file():
        raise ReferencePackBuildError("knowledge_base_input_missing_or_symlink")
    source = requested.resolve()
    text, repairs, status = _read_textual_material(source, budget=budget)
    return {
        "text": text,
        "status": status,
        "repairs": repairs,
        "sha256": f"sha256:{_sha256_file(source)}",
        "material_type": source.suffix.casefold().lstrip("."),
    }


def _text_blocks(value: str, *, limit: int = 16) -> list[str]:
    clean = value.replace("\x00", "")
    clean = re.sub(r"(?m)^\s*#{1,6}\s*", "", clean)
    parts = [_normal(part) for part in re.split(r"\n\s*\n|\n(?=[#\-•])", clean)]
    result: list[str] = []
    for part in parts:
        if not part:
            continue
        for offset in range(0, len(part), 2400):
            piece = part[offset:offset + 2400].strip()
            if piece:
                result.append(piece)
            if len(result) >= limit:
                return result
    return result


def _copy_material(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{destination.name}.", dir=destination.parent)
    os.close(descriptor)
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, destination)
    except Exception:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def _canonical_material_kind(suffix: str) -> str:
    return {
        ".pdf": "user_uploaded_document",
        ".docx": "user_uploaded_document",
        ".pptx": "user_uploaded_presentation",
        ".xlsx": "user_uploaded_spreadsheet",
        ".xls": "user_uploaded_spreadsheet",
    }.get(suffix.casefold(), "user_uploaded_material")


def _source_type_for(suffix: str) -> str:
    return "user_upload" if suffix.casefold() in IMAGE_SUFFIXES else "first_party_internal"


def _claim_type_for(text: str) -> str:
    lowered = text.casefold()
    if any(token in lowered for token in ("价格", "报价", "费用", "price")):
        return "price"
    if any(token in lowered for token in ("活动", "发布", "开业", "event")):
        return "event"
    if any(token in lowered for token in ("限制", "风险", "注意", "不适用", "limit")):
        return "limitation"
    return "knowledge_statement"


def _image_dimensions(path: Path) -> tuple[int | None, int | None]:
    try:
        return validate_material(path)
    except (IntakeSecurityError, OSError):
        return None, None


def _manifest_materials(items: list[dict[str, Any]], *, with_research: bool = False) -> dict[str, Any]:
    payload: dict[str, Any] = {"schema_version": "2.2.0", "items": items}
    if with_research:
        payload["question_research_index_path"] = "research/question_research/index.json"
    return payload


def _portable_members(pack_root: Path) -> list[Path]:
    members: list[Path] = []
    for path in pack_root.rglob("*"):
        if path.is_symlink():
            raise ReferencePackBuildError(
                f"reference_pack_contains_symlink: {path.relative_to(pack_root).as_posix()}"
            )
        if not path.is_file():
            continue
        relative = path.relative_to(pack_root)
        if relative.parts[0] == "workspace_data" or any(part.startswith(".") for part in relative.parts):
            continue
        members.append(path)
    return sorted(members, key=lambda value: value.relative_to(pack_root).as_posix())


def _prepare_portable_path(pack_root: Path, portable_zip: Path | None) -> Path:
    requested = Path(portable_zip).expanduser() if portable_zip is not None else pack_root.with_suffix(".zip")
    if requested.is_symlink():
        raise ReferencePackBuildError("reference_pack_portable_zip_must_not_be_symlink")
    destination = requested.resolve()
    if destination.exists() and destination.is_dir():
        raise ReferencePackBuildError("reference_pack_portable_zip_must_be_file")
    try:
        destination.relative_to(pack_root)
    except ValueError:
        pass
    else:
        raise ReferencePackBuildError("reference_pack_portable_zip_must_be_outside_pack")
    destination.parent.mkdir(parents=True, exist_ok=True)
    return destination


def write_portable_zip(pack_root: Path, portable_zip: Path | None = None) -> tuple[Path, str]:
    """Atomically write a deterministic ZIP containing only execution members."""
    root = Path(pack_root).expanduser().resolve()
    destination = _prepare_portable_path(root, portable_zip)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{destination.name}.", dir=destination.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for path in _portable_members(root):
                name = path.relative_to(root).as_posix()
                info = zipfile.ZipInfo(name, ZIP_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = (stat.S_IFREG | 0o644) << 16
                info.flag_bits |= 0x800
                archive.writestr(info, path.read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
        report = validate_reference_pack(temporary)
        if report.get("status") != "pass":
            raise ReferencePackBuildError(
                "generated_portable_reference_pack_invalid: " + "; ".join(report.get("errors", []))
            )
        digest = _sha256_file(temporary)
        os.replace(temporary, destination)
        return destination, digest
    except Exception:
        try:
            temporary.unlink()
        except FileNotFoundError:
            pass
        raise


def _pack_root_payload(path: Path) -> dict[str, Any]:
    source = Path(path).expanduser().resolve()
    try:
        if source.is_dir():
            return json.loads((source / "reference_pack.json").read_text(encoding="utf-8"))
        with zipfile.ZipFile(source) as archive:
            return json.loads(archive.read("reference_pack.json").decode("utf-8"))
    except (OSError, KeyError, UnicodeDecodeError, json.JSONDecodeError, zipfile.BadZipFile) as exc:
        raise ReferencePackBuildError("reference_pack_manifest_unreadable") from exc


def _copy_pack_members(source: Path, destination: Path) -> None:
    source = Path(source).expanduser().resolve()
    if source.is_dir():
        for path in _portable_members(source):
            relative = path.relative_to(source)
            _atomic_bytes(destination / relative, path.read_bytes())
        return
    try:
        with zipfile.ZipFile(source) as archive:
            infos = validate_zip_infos(
                archive.infolist(), label="reference_pack", budget=None,
                enforce_size_limits=False,
            )
            for info in infos:
                if info.is_dir():
                    continue
                target = destination.joinpath(*PurePosixPath(info.filename).parts)
                with archive.open(info, "r") as incoming:
                    descriptor, temporary_name = tempfile.mkstemp(
                        prefix=f".{target.name}.", dir=target.parent if target.parent.exists() else destination,
                    )
                    os.close(descriptor)
                    temporary = Path(temporary_name)
                    target.parent.mkdir(parents=True, exist_ok=True)
                    try:
                        with temporary.open("wb") as outgoing:
                            shutil.copyfileobj(incoming, outgoing, length=1024 * 1024)
                        os.chmod(temporary, 0o644)
                        os.replace(temporary, target)
                    except Exception:
                        temporary.unlink(missing_ok=True)
                        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile, IntakeSecurityError) as exc:
        raise ReferencePackBuildError(f"reference_pack_copy_failed: {exc}") from exc


def create_next_reference_pack_version(
    *,
    pack: Path,
    output: Path,
    artifacts: Mapping[str, bytes | str | Path] | None = None,
    readiness_updates: Mapping[str, Any] | None = None,
    remove_members: Sequence[str] = (),
    portable_zip: Path | None = None,
) -> dict[str, Any]:
    """Create one immutable next version in the same Reference Pack series."""

    source = Path(pack).expanduser().resolve()
    before = validate_reference_pack(source)
    if before.get("status") != "pass":
        raise ReferencePackBuildError(
            "reference_pack_invalid_before_version_commit: " + "; ".join(before.get("errors") or [])
        )
    old_root = _pack_root_payload(source)
    requested = Path(output).expanduser()
    if requested.is_symlink():
        raise ReferencePackBuildError("reference_pack_output_must_not_be_symlink")
    destination = requested.resolve()
    if destination.exists():
        raise ReferencePackBuildError("reference_pack_output_must_not_exist")
    destination.parent.mkdir(parents=True, exist_ok=True)
    portable_destination = _prepare_portable_path(destination, portable_zip)
    if portable_destination.exists():
        raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
    temporary = Path(tempfile.mkdtemp(prefix=f".{destination.name}.building.", dir=destination.parent))
    descriptor, staged_zip_name = tempfile.mkstemp(
        prefix=f".{portable_destination.name}.building.", dir=portable_destination.parent,
    )
    os.close(descriptor)
    staged_zip = Path(staged_zip_name)
    try:
        _copy_pack_members(source, temporary)
        for member in remove_members:
            if not _safe_relative(member):
                raise ReferencePackBuildError(f"reference_pack_remove_path_unsafe: {member}")
            target = temporary.joinpath(*PurePosixPath(member).parts)
            if target.is_dir():
                shutil.rmtree(target)
            else:
                target.unlink(missing_ok=True)
        for member, value in (artifacts or {}).items():
            if not _safe_relative(member):
                raise ReferencePackBuildError(f"reference_pack_artifact_path_unsafe: {member}")
            if isinstance(value, Path):
                data = value.read_bytes()
            elif isinstance(value, bytes):
                data = value
            else:
                data = str(value).encode("utf-8")
            _atomic_bytes(temporary.joinpath(*PurePosixPath(member).parts), data)
        active_readiness = dict(old_root.get("readiness") or {})
        active_readiness.update(dict(readiness_updates or {}))
        version = int(old_root["pack_version"]) + 1
        root = current_root(
            str(old_root["brand_name"]),
            pack_id=str(old_root["pack_id"]),
            pack_version=version,
            parent_pack_version=int(old_root["pack_version"]),
            created_at=str(old_root["created_at"]),
            updated_at=_now(),
            readiness_value=active_readiness,
        )
        _atomic_json(temporary / "reference_pack.json", root)
        report = validate_reference_pack(temporary)
        if report.get("status") != "pass":
            raise ReferencePackBuildError(
                "reference_pack_invalid_after_version_commit: " + "; ".join(report.get("errors") or [])
            )
        _, zip_digest = write_portable_zip(temporary, staged_zip)
        os.replace(temporary, destination)
        try:
            if portable_destination.exists():
                raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
            os.replace(staged_zip, portable_destination)
        except Exception:
            shutil.rmtree(destination, ignore_errors=True)
            raise
        return {
            "status": "created",
            "builder_version": BUILDER_VERSION,
            "pack_path": str(destination),
            "portable_zip_path": str(portable_destination),
            "portable_zip_sha256": f"sha256:{zip_digest}",
            "pack_id": root["pack_id"],
            "pack_version": version,
            "parent_pack_version": root["parent_pack_version"],
            "brand_name": root["brand_name"],
            "readiness": report["readiness"],
        }
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        staged_zip.unlink(missing_ok=True)
        raise


def upgrade_reference_pack_40_for_market_refresh(*, pack: Path, output: Path) -> dict[str, Any]:
    """Stage Pack 4.0 materials for a 4.1 market refresh.

    The upgrade preserves the immutable series identity, ordinary materials,
    registries and question research. It removes the earlier market strategy
    and P0 because runtime 4.11 must regenerate and reconfirm positioning
    before either can be used again.
    """

    source = Path(pack).expanduser().resolve()
    old_root = _pack_root_payload(source)
    if old_root.get("schema_version") != "4.0" or old_root.get("profile") != "frontmind-content-reference-pack":
        raise ReferencePackBuildError("reference_pack_4_0_upgrade_requires_pack_4_0")
    required = (
        "reference_pack.json", "materials/index.json",
        "registries/source_registry.json", "registries/knowledge_registry.json",
        "registries/claim_registry.json", "registries/image_registry.json",
    )
    requested = Path(output).expanduser()
    if requested.exists() or requested.is_symlink():
        raise ReferencePackBuildError("reference_pack_upgrade_output_must_not_exist")
    destination = requested.resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=f".{destination.name}.upgrading.", dir=destination.parent))
    try:
        _copy_pack_members(source, temporary)
        missing = [member for member in required if not (temporary / member).is_file()]
        if missing:
            raise ReferencePackBuildError("reference_pack_4_0_upgrade_missing_members: " + ", ".join(missing))
        for member in (*BRAND_MARKET_MEMBERS.values(), *STRATEGY_MEMBERS.values(), *P0_MEMBERS.values()):
            target = temporary.joinpath(*PurePosixPath(member).parts)
            if target.is_dir():
                shutil.rmtree(target)
            else:
                target.unlink(missing_ok=True)
        readiness = {
            "materials_ready": True,
            "brand_market_research_ready": False,
            "positioning_ready": False,
            "p0_ready": False,
            "question_ready": dict(old_root.get("readiness", {}).get("question_ready") or {}),
        }
        root = current_root(
            str(old_root.get("brand_name") or ""),
            pack_id=str(old_root.get("pack_id") or ""),
            pack_version=int(old_root.get("pack_version") or 1),
            parent_pack_version=old_root.get("parent_pack_version"),
            created_at=str(old_root.get("created_at") or _now()),
            updated_at=_now(), readiness_value=readiness,
        )
        _atomic_json(temporary / "reference_pack.json", root)
        report = validate_reference_pack(temporary)
        if report.get("status") != "pass" or not report.get("readiness", {}).get("materials_ready"):
            raise ReferencePackBuildError(
                "reference_pack_4_0_upgrade_invalid: " + "; ".join(report.get("errors") or [])
            )
        os.replace(temporary, destination)
        return {
            "status": "staged_for_market_refresh", "builder_version": BUILDER_VERSION,
            "pack_path": str(destination), "pack_id": root["pack_id"],
            "pack_version": root["pack_version"], "brand_name": root["brand_name"],
            "readiness": report["readiness"], "source_schema_version": "4.0",
            "target_schema_version": root["schema_version"], "p0_review_required": True,
        }
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        raise


def _no_usable_error(reports: Sequence[dict[str, Any]]) -> str:
    reasons = [str(item.get("reason")) for item in reports if item.get("status") == "rejected" and item.get("reason")]
    if len(reasons) == 1:
        return reasons[0]
    archive_reason = next((reason for reason in reasons if "archive" in reason), None)
    return archive_reason or reasons[0] if reasons else "knowledge_base_has_no_usable_materials"


def create_reference_pack(
    *, brand_name: str, knowledge_base: Sequence[Path], output: Path,
    portable_zip: Path | None = None,
) -> dict[str, Any]:
    """Create an atomic directory Pack and deterministic portable ZIP."""
    brand = _normal(brand_name)
    if not brand:
        raise ReferencePackBuildError("brand_name_required")
    requested_destination = Path(output).expanduser()
    if requested_destination.is_symlink():
        raise ReferencePackBuildError("reference_pack_output_must_not_be_symlink")
    destination = requested_destination.resolve()
    if destination.exists() and not destination.is_dir():
        raise ReferencePackBuildError("reference_pack_output_must_be_directory")
    if destination.exists() and any(destination.iterdir()):
        raise ReferencePackBuildError("reference_pack_output_not_empty")
    destination.parent.mkdir(parents=True, exist_ok=True)
    portable_destination = _prepare_portable_path(destination, portable_zip)
    if portable_destination.exists():
        raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
    temporary = Path(tempfile.mkdtemp(prefix=f".{destination.name}.building.", dir=destination.parent))
    descriptor, staged_zip_name = tempfile.mkstemp(
        prefix=f".{portable_destination.name}.building.", dir=portable_destination.parent,
    )
    os.close(descriptor)
    staged_zip = Path(staged_zip_name)
    intake_staging = Path(tempfile.mkdtemp(prefix="frontmind-kb-intake-"))
    timestamp = _now()
    repairs: list[dict[str, Any]] = []
    material_items: list[dict[str, Any]] = []
    source_entries: list[dict[str, Any]] = []
    knowledge_units: list[dict[str, Any]] = []
    claims: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    total_extracted_chars = 0
    intake_budget = IntakeBudget()
    candidates, reports = _collect_inputs(
        [Path(item) for item in knowledge_base], intake_staging,
        budget=intake_budget,
    )
    unique: list[MaterialCandidate] = []
    by_hash: dict[str, MaterialCandidate] = {}
    for candidate in candidates:
        canonical = by_hash.get(candidate.sha256)
        if canonical is None:
            by_hash[candidate.sha256] = candidate
            unique.append(candidate)
            continue
        canonical.aliases.extend(alias for alias in candidate.aliases if alias not in canonical.aliases)
        reports[candidate.report_index]["status"] = "skipped"
        reports[candidate.report_index]["reason"] = "duplicate_content"
        reports[candidate.report_index]["aliases"] = list(candidate.aliases)
    if not unique:
        shutil.rmtree(temporary, ignore_errors=True)
        shutil.rmtree(intake_staging, ignore_errors=True)
        raise ReferencePackBuildError(_no_usable_error(reports))
    try:
        for candidate in unique:
            suffix = candidate.path.suffix.casefold()
            digest = candidate.sha256
            original_name = PurePosixPath(
                candidate.source_path.replace("!/", "/")
            ).name
            name = _safe_name(original_name, f"material{suffix}")
            member = f"materials/{digest[:16]}_{name}"
            target = temporary / member
            try:
                raw, repaired_actions, material_status = _read_textual_material(
                    candidate.path, budget=intake_budget,
                )
            except ReferencePackBuildError as exc:
                reports[candidate.report_index]["status"] = "rejected"
                reports[candidate.report_index]["reason"] = str(exc)
                continue
            total_extracted_chars += len(raw)
            if total_extracted_chars > MAX_EXTRACTED_TEXT_CHARS_PER_JOB:
                raise ReferencePackBuildError("knowledge_base_job_text_budget_exceeded")
            if suffix in {".md", ".markdown"}:
                _atomic_text(target, raw)
            else:
                _copy_material(candidate.path, target)
            for repair in repaired_actions:
                repairs.append({"source": candidate.source_path, "repair": repair})
            source_id = f"src_{digest[:16]}"
            material_id = f"material_{digest[:16]}"
            aliases = list(dict.fromkeys(candidate.aliases))
            reports[candidate.report_index]["aliases"] = aliases
            material_items.append({
                "material_id": material_id,
                "path": member,
                "kind": _canonical_material_kind(suffix),
                "usage_status": material_status,
                "source_sha256": f"sha256:{digest}",
                "source_aliases": aliases,
                "original_relative_paths": aliases,
                "material_type": suffix.lstrip("."),
                "publishability_status": "requires_source_review",
                "public_use_permission": "unconfirmed",
                "confidentiality": "unclassified",
            })
            source_entries.append({
                "source_id": source_id,
                "title": Path(aliases[0]).stem or candidate.path.name,
                "publisher": brand,
                "source_type": _source_type_for(suffix),
                "source_origin_class": "first_party_internal",
                "url": None,
                "file_path": member,
                "published_at": None,
                "accessed_at": timestamp,
                "usage_status": material_status,
                "qualification": "用户提交材料；公开使用前按原始材料核对。",
                "notes": "由 Reference Pack Builder 按内容 Hash 登记。",
                "source_hash": f"sha256:{digest}",
                "aliases": aliases,
                "original_relative_paths": aliases,
                "material_type": suffix.lstrip("."),
                "publishability_status": "requires_source_review",
                "public_use_permission": "unconfirmed",
                "confidentiality": "unclassified",
            })
            if suffix in IMAGE_SUFFIXES:
                width, height = _image_dimensions(candidate.path)
                images.append({
                    "image_id": f"img_{digest[:16]}",
                    "file_path": member,
                    "asset_kind": "user_uploaded_image",
                    "usage_status": "usable",
                    "rights_status": "unknown",
                    "allowed_roles": [],
                    "attribution_text": None,
                    "width": width,
                    "height": height,
                    "source_hash": f"sha256:{digest}",
                    "source_ids": [source_id],
                    "original_relative_paths": aliases,
                })
                continue
            for unit_no, block in enumerate(_text_blocks(raw), 1):
                knowledge_id = f"kn_{digest[:12]}_{unit_no:02d}"
                claim_id = f"clm_{digest[:12]}_{unit_no:02d}"
                title = f"{Path(aliases[0]).stem or candidate.path.name}（片段 {unit_no}）"
                knowledge_units.append({
                    "knowledge_id": knowledge_id,
                    "title": title,
                    "kind": "source_derived_text",
                    "content": block,
                    "package_path": member,
                    "source_ids": [source_id],
                    "asset_ids": [],
                    "usage_status": "usable",
                    "qualification": "保留来源片段，不代表额外核验或结论。",
                })
                claims.append({
                    "claim_id": claim_id,
                    "claim_text": block,
                    "claim_type": _claim_type_for(block),
                    "about_entities": [brand],
                    "source_ids": [source_id],
                    "usage_status": "usable",
                    "qualification": "仅可按原始企业材料范围使用；需时效或第三方核验时另行核对。",
                    "verification_status": "needs_verification",
                    "allowed_usage": "internal_only",
                    "as_of": None,
                })
        if not material_items:
            raise ReferencePackBuildError(_no_usable_error(reports))
        identity_seed = json.dumps(
            {
                "brand": brand,
                "materials": sorted(item["source_sha256"] for item in material_items),
            },
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        pack_id = f"rp_{hashlib.sha256(identity_seed).hexdigest()[:16]}"
        _atomic_json(temporary / "reference_pack.json", current_root(
            brand,
            pack_id=pack_id,
            readiness_value={
                "materials_ready": True,
                "brand_market_research_ready": False,
                "positioning_ready": False,
                "p0_ready": False,
                "question_ready": {},
            },
        ))
        _atomic_json(temporary / "materials/index.json", _manifest_materials(material_items))
        _atomic_json(temporary / "registries/knowledge_registry.json", {
            "schema_version": "2.2.0", "registry_type": "knowledge_registry", "knowledge_units": knowledge_units,
        })
        _atomic_json(temporary / "registries/source_registry.json", {
            "schema_version": "2.2.0", "registry_type": "source_registry", "sources": source_entries,
        })
        _atomic_json(temporary / "registries/claim_registry.json", {
            "schema_version": "2.2.0", "registry_type": "claim_registry", "claims": claims,
        })
        _atomic_json(temporary / "registries/image_registry.json", {
            "schema_version": "2.2.0", "registry_type": "image_registry", "images": images,
        })
        report = validate_reference_pack(temporary)
        if report["status"] != "pass" or not report["readiness"]["materials_ready"]:
            raise ReferencePackBuildError("generated_reference_pack_invalid: " + "; ".join(report["errors"]))
        _, zip_digest = write_portable_zip(temporary, staged_zip)
        if destination.exists():
            destination.rmdir()
        os.replace(temporary, destination)
        try:
            if portable_destination.exists():
                raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
            os.replace(staged_zip, portable_destination)
        except Exception:
            shutil.rmtree(destination, ignore_errors=True)
            raise
    except Exception:
        shutil.rmtree(temporary, ignore_errors=True)
        staged_zip.unlink(missing_ok=True)
        raise
    finally:
        shutil.rmtree(intake_staging, ignore_errors=True)
    counts = {status: sum(item["status"] == status for item in reports) for status in ("accepted", "skipped", "rejected")}
    return {
        "status": "created",
        "builder_version": BUILDER_VERSION,
        "pack_path": str(destination),
        "portable_zip_path": str(portable_destination),
        "portable_zip_sha256": f"sha256:{zip_digest}",
        "brand_name": brand,
        "pack_id": pack_id,
        "pack_version": 1,
        "parent_pack_version": None,
        "material_count": len(material_items),
        "knowledge_unit_count": len(knowledge_units),
        "claim_count": len(claims),
        "image_count": len(images),
        "accepted_count": counts["accepted"],
        "skipped_count": counts["skipped"],
        "rejected_count": counts["rejected"],
        "input_items": reports,
        "local_repairs": repairs,
        "readiness": validate_reference_pack(destination)["readiness"],
    }


def _open_directory(path: Path) -> DirectoryView:
    requested = Path(path).expanduser()
    if requested.is_symlink():
        raise ReferencePackBuildError("reference_pack_update_requires_regular_directory")
    candidate = requested.resolve()
    if not candidate.is_dir():
        raise ReferencePackBuildError("reference_pack_update_requires_regular_directory")
    return DirectoryView(candidate)


def validate_reference_pack(path: Path) -> dict[str, Any]:
    """Validate a directory or portable ZIP through the same canonical validator."""
    source = Path(path).expanduser().resolve()
    try:
        if source.is_dir():
            view: Any = DirectoryView(source)
        else:
            from shared.scripts.validate_execution_reference_pack import ZipView
            view = ZipView(source)
        try:
            report = validate_pack(view, execution_compatible=True)
        finally:
            archive = getattr(view, "archive", None)
            if archive is not None:
                archive.close()
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        return {
            "status": "fail", "validator": "v4.11_reference_pack_4.1",
            "errors": [str(exc)], "warnings": [],
            "readiness": {
                "materials_ready": False,
                "brand_market_research_ready": False,
                "positioning_ready": False,
                "p0_ready": False,
                "question_ready": {},
            },
        }
    readiness = report.get("readiness")
    if not isinstance(readiness, dict):
        readiness = {
            "materials_ready": False,
            "brand_market_research_ready": False,
            "positioning_ready": False,
            "p0_ready": False,
            "question_ready": {},
        }
    report["readiness"] = readiness
    report["validator"] = "v4.11_reference_pack_4.1"
    return report


def _load_pack_brand(pack_root: Path) -> str:
    try:
        root = json.loads((pack_root / "reference_pack.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReferencePackBuildError("reference_pack_manifest_unreadable") from exc
    brand = _normal(root.get("brand_name"))
    if not brand:
        raise ReferencePackBuildError("reference_pack_brand_name_required")
    return brand


def _load_material_index(pack_root: Path) -> dict[str, Any]:
    try:
        payload = json.loads((pack_root / "materials/index.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReferencePackBuildError("reference_pack_material_index_unreadable") from exc
    if not isinstance(payload, dict):
        raise ReferencePackBuildError("reference_pack_material_index_invalid")
    return payload


def _read_pack_member(path: Path, member: str) -> bytes:
    source = Path(path).expanduser().resolve()
    try:
        if source.is_dir():
            return (source / member).read_bytes()
        with zipfile.ZipFile(source) as archive:
            return archive.read(member)
    except (OSError, KeyError, zipfile.BadZipFile) as exc:
        raise ReferencePackBuildError(f"reference_pack_member_unreadable: {member}") from exc


def _read_pack_json(path: Path, member: str) -> dict[str, Any]:
    try:
        value = json.loads(_read_pack_member(path, member).decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ReferencePackBuildError(f"reference_pack_json_unreadable: {member}") from exc
    if not isinstance(value, dict):
        raise ReferencePackBuildError(f"reference_pack_json_root_invalid: {member}")
    return value


def add_materials(
    *, pack: Path, materials: Sequence[Path], output: Path,
    portable_zip: Path | None = None,
) -> dict[str, Any]:
    """Append ordinary materials in a new Pack version and reopen positioning.

    Existing question research is preserved. Brand-market research, strategy
    and P0 remain available only in the immutable earlier version; the new
    version removes their canonical members so its readiness truthfully
    returns to materials-only until the existing positioning confirmation flow
    is run again.
    """

    source = Path(pack).expanduser().resolve()
    before = validate_reference_pack(source)
    if before.get("status") != "pass":
        raise ReferencePackBuildError(
            "reference_pack_invalid_before_material_update: "
            + "; ".join(before.get("errors") or [])
        )
    if not materials:
        raise ReferencePackBuildError("reference_pack_material_update_requires_input")
    old_root = _pack_root_payload(source)
    workspace = Path(tempfile.mkdtemp(prefix="frontmind-material-update-"))
    supplemental = workspace / "supplemental"
    try:
        created = create_reference_pack(
            brand_name=str(old_root["brand_name"]),
            knowledge_base=[Path(item) for item in materials],
            output=supplemental,
            portable_zip=workspace / "supplemental.zip",
        )
        del created
        old_index = _read_pack_json(source, "materials/index.json")
        new_index = _read_pack_json(supplemental, "materials/index.json")
        known_hashes = {
            str(item.get("source_sha256") or "")
            for item in old_index.get("items", []) if isinstance(item, dict)
        }
        accepted_materials = [
            item for item in new_index.get("items", [])
            if isinstance(item, dict) and str(item.get("source_sha256") or "") not in known_hashes
        ]
        accepted_paths = {str(item["path"]) for item in accepted_materials}
        accepted_source_ids = {
            f"src_{str(item.get('source_sha256') or '').removeprefix('sha256:')[:16]}"
            for item in accepted_materials
        }
        old_index.setdefault("items", []).extend(accepted_materials)
        artifacts: dict[str, bytes | str | Path] = {
            "materials/index.json": json.dumps(old_index, ensure_ascii=False, indent=2) + "\n",
        }
        for member in accepted_paths:
            artifacts[member] = _read_pack_member(supplemental, member)

        registry_specs = {
            "registries/source_registry.json": ("sources", "source_id"),
            "registries/knowledge_registry.json": ("knowledge_units", "knowledge_id"),
            "registries/claim_registry.json": ("claims", "claim_id"),
            "registries/image_registry.json": ("images", "image_id"),
        }
        for member, (field, identity) in registry_specs.items():
            old_registry = _read_pack_json(source, member)
            new_registry = _read_pack_json(supplemental, member)
            existing_ids = {
                str(item.get(identity)) for item in old_registry.get(field, []) if isinstance(item, dict)
            }
            additions: list[dict[str, Any]] = []
            for item in new_registry.get(field, []):
                if not isinstance(item, dict) or str(item.get(identity)) in existing_ids:
                    continue
                source_ids = set(item.get("source_ids") or [])
                direct_source = str(item.get("source_id") or "")
                file_path = str(item.get("file_path") or item.get("package_path") or "")
                if (
                    direct_source in accepted_source_ids
                    or bool(source_ids & accepted_source_ids)
                    or file_path in accepted_paths
                ):
                    additions.append(item)
            old_registry.setdefault(field, []).extend(additions)
            artifacts[member] = json.dumps(old_registry, ensure_ascii=False, indent=2) + "\n"

        remove_members = [
            *BRAND_MARKET_MEMBERS.values(), *STRATEGY_MEMBERS.values(), *P0_MEMBERS.values(),
        ]
        readiness = {
            "brand_market_research_ready": False,
            "positioning_ready": False,
            "p0_ready": False,
            "question_ready": dict(old_root.get("readiness", {}).get("question_ready") or {}),
        }
        result = create_next_reference_pack_version(
            pack=source, output=Path(output), artifacts=artifacts,
            readiness_updates=readiness, remove_members=remove_members,
            portable_zip=portable_zip,
        )
        result["material_count_added"] = len(accepted_materials)
        result["positioning_reconfirmation_required"] = True
        return result
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


def update_research(
    *, pack: Path, monitoring_answers: Path, source_workbook: Path,
    questions: Sequence[str] = (), output: Path | None = None,
    portable_zip: Path | None = None,
) -> dict[str, Any]:
    """Add question research in a new immutable Reference Pack version."""
    requested_pack = Path(pack).expanduser()
    source_root = requested_pack.resolve()
    before = validate_reference_pack(source_root)
    if before["status"] != "pass":
        raise ReferencePackBuildError("reference_pack_invalid_before_research_update: " + "; ".join(before["errors"]))
    old_root = _pack_root_payload(source_root)
    destination = (
        Path(output).expanduser().resolve()
        if output is not None
        else source_root.with_name(f"{source_root.stem}_v{int(old_root['pack_version']) + 1}")
    )
    if destination.exists() or destination.is_symlink():
        raise ReferencePackBuildError("reference_pack_output_must_not_exist")
    destination.parent.mkdir(parents=True, exist_ok=True)
    portable_destination = _prepare_portable_path(destination, portable_zip)
    if portable_destination.exists():
        raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
    workspace = Path(tempfile.mkdtemp(prefix=f".{destination.name}.research.", dir=destination.parent))
    pack_root = workspace / "pack"
    staged_zip = workspace / "pack.zip"
    try:
        create_next_reference_pack_version(
            pack=source_root,
            output=pack_root,
            readiness_updates={
                "question_ready": dict(old_root.get("readiness", {}).get("question_ready") or {})
            },
            portable_zip=staged_zip,
        )
        brand = _load_pack_brand(pack_root)
        submitted_questions = [
            {"question_text": _normal(item), "source_kind": "project_optimization_target"}
            for item in questions if _normal(item)
        ]
        staging_root = pack_root / "workspace_data/research_ingestion"
        legacy_path = staging_root / "index.json"
        existing_ingestion = research_ingestion.empty_index()
        try:
            ingested, ingestion_actions = research_ingestion.ingest_pairs(
                existing_ingestion,
                [(Path(monitoring_answers), Path(source_workbook))],
                staging_root / "staging", submitted_questions, brand,
            )
            research_ingestion.write_index(legacy_path, ingested)
            existing_question_index = load_question_research_index(pack_root / "research/question_research")
            result = build_from_research_ingestion_index(
                ingested, staging_root / "staging", pack_root / "research/question_research",
                canonical_brand=brand, project_questions=submitted_questions,
                existing_index=existing_question_index,
            )
        except (research_ingestion.ResearchIngestionError, QuestionResearchError, OSError, ValueError) as exc:
            raise ReferencePackBuildError(f"research_update_invalid: {exc}") from exc
        material_index = _load_material_index(pack_root)
        material_index["question_research_index_path"] = "research/question_research/index.json"
        _atomic_json(pack_root / "materials/index.json", material_index)
        shutil.rmtree(pack_root / "workspace_data", ignore_errors=True)
        probe = validate_reference_pack(pack_root)
        root = _pack_root_payload(pack_root)
        root["readiness"]["question_ready"] = probe.get("readiness", {}).get("question_ready", {})
        root["updated_at"] = _now()
        _atomic_json(pack_root / "reference_pack.json", root)
        after = validate_reference_pack(pack_root)
        if after["status"] != "pass":
            raise ReferencePackBuildError(
                "reference_pack_invalid_after_research_update: " + "; ".join(after["errors"])
            )
        staged_zip.unlink(missing_ok=True)
        _, zip_digest = write_portable_zip(pack_root, staged_zip)
        os.replace(pack_root, destination)
        try:
            if portable_destination.exists():
                raise ReferencePackBuildError("reference_pack_portable_zip_must_not_exist")
            os.replace(staged_zip, portable_destination)
        except Exception:
            shutil.rmtree(destination, ignore_errors=True)
            raise
        return {
            "status": "updated",
            "builder_version": BUILDER_VERSION,
            "pack_path": str(destination),
            "portable_zip_path": str(portable_destination),
            "portable_zip_sha256": f"sha256:{zip_digest}",
            "brand_name": brand,
            "pack_id": root["pack_id"],
            "pack_version": root["pack_version"],
            "parent_pack_version": root["parent_pack_version"],
            "ingestion_actions": ingestion_actions,
            "research_actions": list(result.actions),
            "question_uids": list(result.question_uids),
            "readiness": after["readiness"],
        }
    finally:
        shutil.rmtree(workspace, ignore_errors=True)


__all__ = [
    "ACCEPTED_SUFFIXES", "ARCHIVE_SUFFIXES", "BUILDER_VERSION",
    "MAX_COMPRESSION_RATIO", "MAX_INPUT_BYTES", "MAX_INPUT_FILES", "MAX_TOTAL_BYTES",
    "ReferencePackBuildError", "add_materials", "create_next_reference_pack_version",
    "create_reference_pack", "is_knowledge_base_input", "update_research",
    "upgrade_reference_pack_40_for_market_refresh", "validate_reference_pack", "write_portable_zip",
]
